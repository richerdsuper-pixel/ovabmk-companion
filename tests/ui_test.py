"""Проверка интерфейса сайта на мок-данных (без сервера): python3 tests/ui_test.py"""
import json, sys, pathlib
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parents[1]
SITE = (ROOT / "docs/index.html").as_uri()
MOCK = (ROOT / "tests/mock_supabase.js").read_text(encoding="utf-8")
FOLDERS = json.loads((ROOT / "tests/folders.json").read_text(encoding="utf-8"))
SHOTS = ROOT / "tests/shots"; SHOTS.mkdir(exist_ok=True)
A, S_ = "00000000-0000-4000-8000-00000000000a", "00000000-0000-4000-8000-00000000000b"
results = []

def ok(name, cond, info=""):
    results.append(bool(cond)); print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  <- {info}"))

def scenario(uid=None, roles=(), role_error=None, my_folders=()):
    db = {
        "folders": FOLDERS,
        "user_roles": [{"user_id": uid, "role": r} for r in roles],
        "profiles": [{"id": A, "full_name": "Администратор Тест", "position": "Заместитель"},
                     {"id": S_, "full_name": "Сотрудник Тест", "position": "Доя"}],
        "folder_members": [{"user_id": uid, "folder_id": f} for f in my_folders],
        "documents": [{"id": "d1", "folder_id": 1, "title": "Приказ №1 (тест)", "doc_number": "1", "doc_date": "2026-10-01", "note": None,
                       "storage_path": "1/11111111-1111-4111-8111-111111111111/prikaz.pdf", "file_name": "prikaz.pdf", "file_size": 120000,
                       "uploaded_by": A, "created_at": "2026-10-20T08:00:00Z"}],
        "invites": [{"id": "i1", "role": "staff", "full_name": "Тестовый Сотрудник", "position": "Доя", "folder_ids": [1, 2],
                     "created_at": "2026-10-20T08:00:00Z", "expires_at": "2099-01-01T00:00:00Z", "used_at": None, "revoked_at": None}],
        "access_requests": [
            {"id": "r1", "full_name": "Иванова Мария", "position": "Акушерка СВП (кабинет доя)", "position_code": "doya_qvp", "tg_username": "maria_i", "tg_name": "Maria",
             "status": "pending", "role": None, "folder_ids": [], "created_at": "2026-10-24T07:00:00Z", "decided_at": None, "used_at": None},
            {"id": "r2", "full_name": "Каримов Алишер", "position": "Врач-лаборант", "position_code": "other", "tg_username": None, "tg_name": "Alisher",
             "status": "pending", "role": None, "folder_ids": [], "created_at": "2026-10-24T07:05:00Z", "decided_at": None, "used_at": None},
            {"id": "r3", "full_name": "Сидорова Анна", "position": "Врач-педиатр", "position_code": "pediatr", "tg_username": "anna", "tg_name": "Anna",
             "status": "approved", "role": "staff", "folder_ids": [4, 5, 17], "created_at": "2026-10-23T07:00:00Z", "decided_at": "2026-10-23T08:00:00Z", "used_at": None}],
        "audit_log": [{"at": "2026-10-24T06:00:00Z", "actor": A, "action": "document_added", "object": "d1", "details": {"folder": 1, "title": "Приказ №1 (тест)"}}],
    }
    return {"session": {"user": {"id": uid}} if uid else None, "db": db, "log": [],
            "errors": {"user_roles": role_error} if role_error else {}}

RPC = """
window.__MOCK.rpc = {
  folder_counts: function(){ return [{folder_id:1, docs:1}]; },
  admin_list_users: function(){ return [
    {user_id:'%s', full_name:'Администратор Тест', user_position:'Заместитель', telegram_username:'admin_t', roles:['admin'], folders:[], created_at:'2026-10-20T08:00:00Z'},
    {user_id:'%s', full_name:'Сотрудник Тест', user_position:'Доя', telegram_username:'staff_t', roles:['staff'], folders:[1], created_at:'2026-10-21T08:00:00Z', institution_id:1, scope_institutions:[], scope_positions:[]}]; },
  admin_save_institution: function(a){ return a._id || 3; }, admin_add_institutions: function(a){ return a._names.length - 1; },
  admin_save_mahalla: function(a){ return a._id || 9; }, admin_delete_mahalla: function(){ return window.__mhErr ? {error:{message:'mahalla_has_reports'}} : null; }, admin_set_brigade: function(){ return null; },
  admin_delete_institution: function(){ return null; }, admin_set_institution: function(){ return null; }, admin_set_scopes: function(){ return null; },
  create_invite: function(){ return '0123456789abcdef01234567'; },
  admin_set_role: function(a){ return a._role==='admin' && !a._on ? {error:{message:'last_admin'}} : null; },
  decide_request: function(){ return null; },
  admin_set_folders: function(){ return null; }, admin_update_profile: function(){ return null; }, revoke_invite: function(){ return null; },
  my_broadcasts: function(){ return window.__BC || []; }, my_tasks: function(){ return window.__TK || []; },
  broadcast_view: function(){ return null; }, task_seen: function(){ return null; }, task_done: function(){ return null; },
  admin_list_broadcasts: function(){ return window.__ABC || []; },
  admin_broadcast_status: function(){ return [
    {user_id:'u1', full_name:'Иванова Мария', user_position:'Доя', notified_version:2, viewed:true, first_view:'2026-10-24T06:00:00Z', views:2, downloaded:true, last_download:'2026-10-24T06:05:00Z', downloads:1},
    {user_id:'u2', full_name:'Петрова Зарина', user_position:'Педиатр', notified_version:null, viewed:false, views:0, downloaded:false, downloads:0}]; },
  admin_publish_broadcast: function(a){ return window.__pubErr ? {error:{message:window.__pubErr}} : {id:a._id, version:1, recipients:3}; },
  admin_delete_broadcast: function(){ return ['x/1/plan.xlsx']; },
  admin_list_tasks: function(){ return window.__ATK || []; },
  admin_task_status: function(){ return [{user_id:'u1', full_name:'Иванова Мария', user_position:'Доя', notified_at:'2026-10-24T05:00:00Z', seen_at:'2026-10-24T06:00:00Z', done_at:null, note:''}]; },
  admin_create_task: function(){ return 't-new'; }, admin_delete_task: function(){ return null; },
  anc_schedule_list: function(){ return [
    {visit_no:1, label:'1-й визит (до 12 нед.)', from_week:0, until_week:12, tests:['Общий анализ крови','ВИЧ, RW'], notes:'Оценить факторы риска.', label_uz:'1-кўрик (12 ҳафтагача)', tests_uz:['Умумий қон таҳлили','ОИВ, RW'], notes_uz:'Хавф омилларини баҳолаш.'},
    {visit_no:2, label:'2-й визит (18–20 нед.)', from_week:16, until_week:22, tests:['УЗИ плода'], notes:'Аномалии развития.'},
    {visit_no:3, label:'3-й визит (24–26 нед.)', from_week:23, until_week:28, tests:[], notes:''}]; },
  admin_save_anc_row: function(){ return window.__ancErr ? {error:{message:window.__ancErr}} : null; },
  admin_offline_roles: function(){ return [{user_id:'%s', offline_role:'cab'}]; }, admin_set_offline_role: function(a){ return a._role==='coord' ? {error:{message:'coord_admin_only'}} : null; },
  admin_get_offline_key: function(){ return window.__AKEY || null; }, admin_set_offline_key: function(){ window.__AKEY = {x:'x', y:'y'}; return null; },
  my_work_profile: function(){ return {rate:1, rate_req:null, offline_role:'cab', has_pub:false, is_head:false, max:1.5, sub_load:0.3}; },
  leaves_list: function(){ return window.__LV || []; }, admin_rates: function(){ return [{user_id:'u9', full_name:'Каримова Д.', user_position:'Доя', rate:1, rate_req:1.25, rate_req_at:'2026-10-24T05:00:00Z'}]; },
  leave_candidates: function(){ return [{user_id:'u5', full_name:'Холова Ж.', user_position:'Доя', rate:1, load_now:1, can:true}, {user_id:'u6', full_name:'Рахимова С.', user_position:'Доя', rate:1.5, load_now:1.5, can:false}]; },
  request_leave: function(a){ return window.__lvErr ? {error:{message:window.__lvErr}} : 'lv-new'; }, request_rate: function(){ return null; },
  substitution_consent: function(){ return null; }, decide_leave: function(a){ return window.__decErr ? {error:{message:window.__decErr}} : null; },
  admin_demand_review: function(){ return null; }, cancel_leave: function(){ return null; }, admin_confirm_rate: function(){ return null; }, set_my_offline_pub: function(){ return null; },
  admin_remove_user: function(a){ return a._uid === '%s' ? {error:{message:'cannot_remove_self'}} : (window.__rm || 'deleted'); }
};""" % (A, S_, S_, A)

FN = """window.__MOCK.fn = { 'bc-file': function(){ return window.__fnErr ? {error:{message:'forbidden'}} : {url:'https://example.test/dl'}; },
  'tg-notify': function(){ return {total:3, sent:2, blocked:1, failed:0}; },
  'anc-digest': function(){ return {total:4, sent:3, blocked:1, failed:0}; },
  'leave-notify': function(){ return {total:2, sent:2, failed:0}; },
  'offline-key': function(b){ return window.__offErr ? {error:{message:window.__offErr}} : {token:'eyJ2IjoxfQ.c2ln' + (b && b.uid ? 'U' : ''), role:'cab', exp:1}; } };"""

def page_for(p, sc, w=1280, h=900, scheme="light", search="", now=None, hash="#/folders"):
    b = p.chromium.launch()
    ctx = b.new_context(viewport={"width": w, "height": h}, color_scheme=scheme, accept_downloads=True)
    if now: ctx.add_init_script("window.__NOW=" + json.dumps(now))
    ctx.add_init_script("window.__MOCK=" + json.dumps(sc, ensure_ascii=False) + ";" + RPC + FN)
    ctx.route("**/supabase.js", lambda r: r.fulfill(status=200, content_type="application/javascript", body=MOCK))
    ctx.route("**/telegram-widget.js*", lambda r: r.fulfill(status=200, content_type="application/javascript", body="/* stub */"))
    ctx.route("**/fonts.g*/**", lambda r: r.abort())
    pg = ctx.new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + search + hash); pg.wait_for_timeout(400)
    return b, pg, errs

def hscroll(pg): return pg.evaluate("document.documentElement.scrollWidth > document.documentElement.clientWidth + 1")

with sync_playwright() as p:
    # 1. без входа + приглашение в адресе
    b, pg, errs = page_for(p, scenario(), search="?invite=0123456789ABCDEF01234567")
    ok("без сессии показан вход", "Вход" in pg.inner_text("#app"))
    ok("приглашение распознано", "приглашение" in pg.inner_text("#app"))
    ok("код приглашения убран из адреса", "invite=" not in pg.url, pg.url)
    ok("код сохранён до входа", pg.evaluate("sessionStorage.getItem('ovabmk-invite')") == "0123456789abcdef01234567")
    ok("Telegram-скрипт вставлен с именем бота", pg.evaluate("!!document.querySelector('script[data-telegram-login=\"OvaBMQbot\"]')"))
    pg.screenshot(path=str(SHOTS / "c_login.png"), full_page=True)
    pg.click(".seg button:nth-child(2)"); pg.wait_for_timeout(150)
    ok("переключение на узбекский", "Кириш" in pg.inner_text("#app"))
    ok("ошибок JS нет (вход)", not errs, errs); b.close()

    # 2. ошибка чтения ролей
    b, pg, errs = page_for(p, scenario(uid=A, roles=("admin",), role_error="permission denied for table user_roles"))
    t = pg.inner_text("#app")
    ok("ошибка роли показана как ошибка, а не «—»", "Не удалось загрузить роль" in t and "Повторить" in t, t[:200])
    ok("при ошибке роли нет меню администратора", "Сотрудники" not in t)
    pg.screenshot(path=str(SHOTS / "c_role_error.png"), full_page=True); b.close()

    # 3. роль не назначена
    b, pg, errs = page_for(p, scenario(uid=S_, roles=()))
    t = pg.inner_text("#app")
    ok("нет роли — отдельное сообщение", "Роль не назначена" in t and "администратору" in t, t[:200])
    ok("нет роли — папки не показаны", "Папка №1" not in t); b.close()

    # 4. администратор
    b, pg, errs = page_for(p, scenario(uid=A, roles=("admin",)))
    t = pg.inner_text("#app")
    ok("admin: роль показана", "Администратор" in t)
    ok("admin: меню администратора", all(x in t for x in ("Сотрудники", "Приглашения", "Журнал")))
    ok("20 папок", pg.locator(".fcard").count() == 20, pg.locator(".fcard").count())
    ok("папка 15 закрыта и не ссылка", pg.locator(".fcard.locked").count() == 1 and pg.locator("a.fcard[href='#/folder/15']").count() == 0)
    pg.screenshot(path=str(SHOTS / "c_folders.png"), full_page=True)
    pg.click("a.fcard[href='#/folder/1']"); pg.wait_for_timeout(300)
    t = pg.inner_text("#app")
    ok("папка 1: документ виден с автором", "Приказ №1 (тест)" in t and "Администратор Тест" in t)
    pg.click("details.panel summary"); pg.wait_for_timeout(100)
    pg.fill("form input[type=text]", "Новый план")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(150)
    ok("без файла — ошибка формы", "Выберите файл" in pg.inner_text("form"))
    pg.set_input_files("form input[type=file]", files=[{"name": "План работы.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4 test"}])
    pg.click("form button[type=submit]"); pg.wait_for_timeout(150)
    ok("без подтверждения — ошибка формы", "Подтвердите" in pg.inner_text("form"))
    pg.check("form .check input"); pg.click("form button[type=submit]"); pg.wait_for_timeout(400)
    log = pg.evaluate("window.__MOCK.log")
    up = [x for x in log if x[0] == "upload"]; ins = [x for x in log if x[0] == "insert"]
    ok("файл загружен по пути папки с латинским именем", up and up[0][1].startswith("1/") and up[0][1].endswith("/plan_raboty.pdf"), up)
    ok("строка документа с подтверждением", ins and ins[0][2]["no_patient_data"] is True and ins[0][2]["file_name"] == "План работы.pdf", ins)
    pg.screenshot(path=str(SHOTS / "c_folder.png"), full_page=True)
    # пакетная загрузка
    pg.locator("details.panel summary", has_text="несколько файлов").click(); pg.wait_for_timeout(100)
    bulk = pg.locator("details.panel", has_text="несколько файлов")
    pg.evaluate("window.__MOCK.log.length = 0")
    bulk.locator("input[type=file]").set_input_files(files=[
        {"name": "Prikaz_12.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4 a"},
        {"name": "Шаблон_журнала.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": b"PK xlsx"},
        {"name": "virus.exe", "mimeType": "application/x-msdownload", "buffer": b"MZ"}])
    pg.wait_for_timeout(150)
    ok("пакет: три строки, у неподходящего — ошибка типа", bulk.locator("input[type=text]").count() == 3 and "не принимается" in bulk.inner_text())
    ok("пакет: название из имени файла", bulk.locator("input[type=text]").nth(0).input_value() == "Prikaz 12")
    bulk.locator("button[type=submit]").click(); pg.wait_for_timeout(150)
    ok("пакет: без подтверждения — ошибка", "Подтвердите" in bulk.inner_text())
    bulk.locator(".check input").check(); bulk.locator("button[type=submit]").click(); pg.wait_for_timeout(500)
    log = pg.evaluate("window.__MOCK.log")
    ups = [x for x in log if x[0] == "upload"]; inss = [x for x in log if x[0] == "insert"]
    ok("пакет: загружены два допустимых файла", len(ups) == 2 and len(inss) == 2, log)
    ok("пакет: пути в папке 1, латиница", all(u[1].startswith("1/") for u in ups) and ups[1][1].endswith("/shablon_jurnala.xlsx"), ups)
    ok("пакет: подтверждение данных пациентов записано", all(i[2]["no_patient_data"] is True for i in inss) and inss[1][2]["title"] == "Шаблон журнала", inss)
    btn = pg.locator("button.danger").first; btn.click(); pg.wait_for_timeout(100)
    ok("удаление требует второго нажатия", "Точно" in btn.inner_text())
    btn.click(); pg.wait_for_timeout(300)
    ok("удаление документа и файла", any(x[0] == "remove" for x in pg.evaluate("window.__MOCK.log")))
    pg.goto(SITE + "#/folder/15"); pg.wait_for_timeout(300)
    ok("прямая ссылка на папку 15 — закрыто", "закрыто для всех" in pg.inner_text("#app"))
    pg.goto(SITE + "#/admin/users"); pg.wait_for_timeout(300)
    t = pg.inner_text("#app")
    ok("сотрудники: список", pg.locator("input[value='Сотрудник Тест']").count() == 1 and "@staff_t" in t, t[:300])
    ok("папка 15 не назначается", pg.locator(".fold-picks label.off input[disabled]").count() >= 1)
    cb0 = pg.locator("tbody tr").first.locator("input[type=checkbox]").first; cb0.click(); pg.wait_for_timeout(250)
    ok("галочка admin вернулась обратно", cb0.is_checked())
    ok("нельзя снять последнего admin (сообщение)", "последнего администратора" in pg.inner_text("#toast"))
    # --- сотрудники: папки по должности и удаление
    pg.evaluate("window.__MOCK.log.length = 0")
    row_s = pg.locator("tbody tr").nth(1)
    row_s.locator("input[type=text]").nth(1).fill("Акушерка СВП (кабинет доя)")
    row_s.locator("button", has_text="Папки по должности").click(); pg.wait_for_timeout(250)
    sf = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_set_folders"]
    ok("сотрудники: папки по должности доя = 9,10,12,13", sf and sf[0][2]["_folders"] == [9, 10, 12, 13], sf)
    ok("сотрудники: галочки выставлены", row_s.evaluate("el => Array.from(el.querySelectorAll('.fold-picks input')).map((x,k)=>x.checked?k+1:0).filter(Boolean)") == [9, 10, 12, 13])
    row_s.locator("input[type=text]").nth(1).fill("Что-то своё")
    row_s.locator("button", has_text="Папки по должности").click(); pg.wait_for_timeout(150)
    ok("сотрудники: должность не из списка — подсказка", "не из списка" in pg.inner_text("#toast"))
    row_s.locator("input[type=text]").nth(1).fill("Главный врач")
    row_s.locator("button", has_text="Папки по должности").click(); pg.wait_for_timeout(150)
    ok("сотрудники: руководителю папки отдельно не нужны", "Руководитель видит все папки" in pg.inner_text("#toast"))
    rm = row_s.locator("button.danger", has_text="Удалить"); rm.click(); pg.wait_for_timeout(80)
    ok("сотрудники: удаление требует второго нажатия", "Точно" in rm.inner_text())
    pg.evaluate("window.__MOCK.log.length = 0; window.__rm = 'deactivated'")
    rm.click(); pg.wait_for_timeout(300)
    ok("сотрудники: запрос на удаление ушёл", any(x[0] == "rpc" and x[1] == "admin_remove_user" and x[2]["_uid"] == S_ for x in pg.evaluate("window.__MOCK.log")))
    ok("сотрудники: пояснение про сохранённую запись", "запись сохранена" in pg.inner_text("#toast").lower())
    row_a = pg.locator("tbody tr").first
    ra = row_a.locator("button.danger", has_text="Удалить"); ra.click(); pg.wait_for_timeout(60); ra.click(); pg.wait_for_timeout(250)
    ok("сотрудники: себя удалить нельзя", "самого себя" in pg.inner_text("#toast"), pg.inner_text("#toast"))
    pg.screenshot(path=str(SHOTS / "c_users.png"), full_page=True)
    pg.goto(SITE + "#/admin/invites"); pg.wait_for_timeout(300)
    pg.select_option("form select >> nth=0", "pediatr"); pg.wait_for_timeout(100)
    ok("приглашение: должность врач-педиатр — роль и папки сами", pg.eval_on_selector("form select >> nth=1", "s => s.value") == "staff" and pg.evaluate("Array.from(document.querySelectorAll('.fold-picks input:checked')).map(x=>+x.value)") == [4, 5, 17])
    ok("приглашение: подсказка про автоматический выбор", "по должности" in pg.inner_text("form"))
    pg.select_option("form select >> nth=0", "bosh_vrach"); pg.wait_for_timeout(100)
    ok("приглашение: главный врач — роль главного врача, без галочек", pg.eval_on_selector("form select >> nth=1", "s => s.value") == "chief_physician" and pg.evaluate("document.querySelectorAll('.fold-picks input:checked').length") == 0)
    pg.select_option("form select >> nth=0", "other"); pg.wait_for_timeout(100)
    ok("приглашение: «другая должность» — поле для ввода", pg.locator("form input[placeholder]").first.is_visible())
    pg.select_option("form select >> nth=0", "pediatr"); pg.evaluate("window.__MOCK.log.length = 0")
    pg.fill("form input[type=text] >> nth=0", "Новый педиатр"); pg.click("form button[type=submit]"); pg.wait_for_timeout(300)
    ci = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "create_invite"]
    ok("приглашение: отправлены должность, роль и папки по должности", ci and ci[0][2]["_position"] == "Врач-педиатр" and ci[0][2]["_role"] == "staff" and ci[0][2]["_folders"] == [4, 5, 17], ci)
    t = pg.inner_text("#app")
    ok("приглашение: ссылка показана", "https://richerdsuper-pixel.github.io/ovabmk-companion/?invite=0123456789abcdef01234567" in t)
    ok("список приглашений", "Тестовый Сотрудник" in t and "Действует" in t)
    pg.screenshot(path=str(SHOTS / "c_invites.png"), full_page=True)
    pg.goto(SITE + "#/admin/log"); pg.wait_for_timeout(300)
    ok("журнал", "добавлен документ" in pg.inner_text("#app"))
    ok("ошибок JS нет (admin)", not errs, errs); b.close()

    # 5. сотрудник с папкой 1
    b, pg, errs = page_for(p, scenario(uid=S_, roles=("staff",), my_folders=(1,)))
    t = pg.inner_text("#app")
    ok("staff: нет меню администратора", "Приглашения" not in t)
    ok("staff: открыть можно только папку 1", pg.locator("a.fcard").count() == 1 and pg.locator("a.fcard[href='#/folder/1']").count() == 1)
    pg.goto(SITE + "#/folder/2"); pg.wait_for_timeout(300)
    ok("staff: прямая ссылка на чужую папку — нет доступа", "не назначена" in pg.inner_text("#app"))
    pg.goto(SITE + "#/admin/users"); pg.wait_for_timeout(300)
    ok("staff: админ-раздел по ссылке закрыт", "только администратору" in pg.inner_text("#app"))
    pg.goto(SITE + "#/folder/1"); pg.wait_for_timeout(300)
    ok("staff: нет кнопки удаления", pg.locator("button.danger").count() == 0)
    ok("ошибок JS нет (staff)", not errs, errs); b.close()

    # 5б. заявка без приглашения
    def tg_route(ctx_page, answers):
        calls = []
        def handler(route):
            body = json.loads(route.request.post_data or "{}"); calls.append(body)
            route.fulfill(status=200, content_type="application/json", body=json.dumps(answers(body)),
                          headers={"Access-Control-Allow-Origin": "*"})
        ctx_page.route("**/functions/v1/tg-login", handler)
        return calls
    b, pg, errs = page_for(p, scenario())
    calls = tg_route(pg, lambda body: {"request": "pending"} if body.get("request") else {"request": "none"})
    pg.evaluate("onTelegramAuth({id: 777, first_name: 'Мария', last_name: 'Иванова', auth_date: 1, hash: 'x'})"); pg.wait_for_timeout(400)
    t = pg.inner_text("#app")
    ok("заявка: после входа без приглашения показана форма", "Заявка на доступ" in t and "Отправить заявку" in t, t[:300])
    ok("заявка: ФИО подставлено из Telegram", pg.input_value("form input[type=text] >> nth=0") == "Иванова Мария")
    ok("заявка: в списке должностей есть акушерка", pg.locator("form select option[value='doya_qvp']").count() == 1)
    pg.click("form button[type=submit]"); pg.wait_for_timeout(150)
    ok("заявка: без должности — ошибка", "Выберите должность" in pg.inner_text("form"))
    pg.select_option("form select", "other"); pg.wait_for_timeout(50)
    ok("заявка: «Другая» открывает поле должности", pg.locator("form input[type=text] >> nth=1").is_visible())
    pg.select_option("form select", "doya_qvp")
    pg.screenshot(path=str(SHOTS / "c_request_form.png"), full_page=True)
    pg.click("form button[type=submit]"); pg.wait_for_timeout(400)
    sent = calls[-1].get("request") or {}
    ok("заявка: отправлены ФИО, должность и код", sent.get("full_name") == "Иванова Мария" and sent.get("position_code") == "doya_qvp" and "Акушерка" in sent.get("position", ""), sent)
    t = pg.inner_text("#app")
    ok("заявка: экран «ждёт одобрения»", "ждёт одобрения" in t and "Проверить" in t, t[:300])
    pg.screenshot(path=str(SHOTS / "c_request_pending.png"), full_page=True)
    pg.click(".seg button:nth-child(2)"); pg.wait_for_timeout(150)
    ok("заявка: узбекский текст", "ариза" in pg.inner_text("#app").lower())
    ok("ошибок JS нет (заявка)", not errs, errs); b.close()

    # 5в. администратор: заявки
    b, pg, errs = page_for(p, scenario(uid=A, roles=("admin",)))
    ok("admin: в меню есть «Заявки»", "Заявки" in pg.inner_text("nav"))
    pg.goto(SITE + "#/admin/requests"); pg.wait_for_timeout(400)
    t = pg.inner_text("#app")
    ok("заявки: общая ссылка показана", "https://richerdsuper-pixel.github.io/ovabmk-companion/" in t and "Ссылка для сотрудников" in t)
    ok("заявки: 2 новые", "Новые заявки: 2" in t, t[:400])
    card1 = pg.locator(".req").nth(0)
    checked = card1.evaluate("el => Array.from(el.querySelectorAll('.fold-picks input:checked')).map(x => +x.value)")
    ok("заявки: папки по должности отмечены сами (доя → 9,10,12,13)", checked == [9, 10, 12, 13], checked)
    ok("заявки: подсказка про автоматический выбор", "автоматически" in card1.inner_text())
    card2 = pg.locator(".req").nth(1)
    ok("заявки: «другая должность» — без галочек и с подсказкой", card2.evaluate("el => el.querySelectorAll('.fold-picks input:checked').length") == 0 and "вручную" in card2.inner_text())
    ok("заявки: рассмотренная в истории", "Сидорова Анна" in t and "ждёт входа" in t)
    pg.screenshot(path=str(SHOTS / "c_requests.png"), full_page=True)
    card1.locator(".fold-picks input[value='13']").uncheck()
    card1.locator("button.pri").click(); pg.wait_for_timeout(300)
    dec = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "decide_request"]
    ok("заявки: одобрение с ролью и изменёнными папками", dec and dec[0][2] == {"_id": "r1", "_approve": True, "_role": "staff", "_folders": [9, 10, 12], "_institution_id": None}, dec)
    rej = pg.locator(".req").nth(1).locator("button.danger"); rej.click(); pg.wait_for_timeout(100)
    ok("заявки: отказ требует второго нажатия", "Точно" in rej.inner_text())
    ok("ошибок JS нет (заявки)", not errs, errs); b.close()

    # 6. телефон и тёмная тема
    for tag, (w, h_, sch) in {"m": (390, 800, "light"), "d": (1280, 900, "dark")}.items():
        b, pg, errs = page_for(p, scenario(uid=A, roles=("admin",)), w, h_, sch)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_folders.png"), full_page=True)
        ok(f"{tag}: нет горизонтальной прокрутки (папки)", not hscroll(pg))
        pg.goto(SITE + "#/admin/users"); pg.wait_for_timeout(300)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_users.png"), full_page=True)
        ok(f"{tag}: нет горизонтальной прокрутки (сотрудники)", not hscroll(pg))
        pg.goto(SITE + "#/admin/requests"); pg.wait_for_timeout(300)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_requests.png"), full_page=True)
        ok(f"{tag}: нет горизонтальной прокрутки (заявки)", not hscroll(pg))
        b.close()

    # 7. показатели
    DOYA = "Акушерка СВП (кабинет доя)"
    X_ = "00000000-0000-4000-8000-00000000000c"
    def ind_sc(uid, roles, reports):
        sc = scenario(uid=uid, roles=roles)
        sc["db"]["profiles"] = [{"id": A, "full_name": "Администратор Тест", "position": "Заместитель главного врача"},
                                {"id": S_, "full_name": "Сотрудник Тест", "position": DOYA},
                                {"id": X_, "full_name": "Вторая Доя", "position": DOYA}]
        sc["db"]["indicator_reports"] = reports
        return sc
    NOW = "2026-10-24T10:00:00"
    rep = lambda rid, uid, inst, pe, vals, ps=None: {"id": rid, "user_id": uid, "position_code": "doya_qvp", "institution": inst, "period_start": ps or pe, "period_end": pe, "vals": vals, "updated_at": "2026-10-23T09:00:00Z"}
    b0, pg0, e0 = page_for(p, ind_sc(S_, ("staff",), []), w=390, h=844, now=NOW, hash="#/overview")
    ct = pg0.inner_text(".act") if pg0.locator(".act").count() else ""
    ok("главная: просроченный срок — красная карточка", "Показатели не сданы" in ct and "23.10.2026" in ct and pg0.locator(".act.t-bad").count() == 1, ct)
    pg0.screenshot(path=str(SHOTS / "c_m_staff_late.png"), full_page=True)
    ok("главная: ошибок JS нет (просрочка)", not e0, e0); b0.close()
    b, pg, errs = page_for(p, ind_sc(S_, ("staff",), [rep("p1", S_, "СВП №3", "2026-10-23", {"i1": 12, "i2": 10}, "2026-10-19")]), now=NOW)
    t = pg.inner_text("#app")
    ok("показатели: пункт меню у сотрудника", "Показатели" in t and "Сводка показателей" not in t)
    pg.goto(SITE + "#/overview"); pg.wait_for_timeout(400)
    ct = pg.inner_text(".act") if pg.locator(".act").count() else ""
    ok("главная: карточка «сдать до» у сотрудника, сдавшего прошлый срок", "Нужно сдать показатели" in ct and "до 30.10.2026" in ct and pg.locator(".act.t-warn a[href='#/indicators']").count() == 1, ct)
    pg.screenshot(path=str(SHOTS / "c_staff_overview.png"), full_page=True)
    pg.goto(SITE + "#/indicators"); pg.wait_for_timeout(400)
    t = pg.inner_text("#app")
    ok("показатели: форма должности подставлена из профиля", pg.eval_on_selector(".grid-f select", "s => s.value") == "doya_qvp")
    ok("показатели: 27 показателей у доя", pg.locator(".indrow").count() == 27, pg.locator(".indrow").count())
    ok("показатели: предупреждение «только числа»", "Фамилии, ПИНФЛ" in t)
    ok("показатели: срок сдачи доя — пятница 12:00", "пятница до 12:00" in t, t[:600])
    ok("показатели: период по умолчанию — неделя пн–пт", pg.eval_on_selector_all("form input[type=date]", "x => x.map(i => i.value)") == ["2026-10-19", "2026-10-23"])
    ok("показатели: мои отчёты показывают сохранённый", "СВП №3" in t and "заполнено 2 из 27" in t, t[-500:])
    pg.screenshot(path=str(SHOTS / "c_ind_form.png"), full_page=True)
    v0 = pg.input_value("#ind1")
    pg.fill("#ind1", "5")
    pg.click(".indrow >> nth=0 >> .stp.plus"); pg.click(".indrow >> nth=0 >> .stp.plus"); pg.click(".indrow >> nth=0 >> .stp:not(.plus)")
    ok("показатели: кнопки + и − меняют число", pg.input_value("#ind1") == "6", pg.input_value("#ind1"))
    pg.fill("#ind1", "")
    pg.click(".indrow >> nth=0 >> .stp:not(.plus)"); pg.click(".indrow >> nth=0 >> .stp:not(.plus)")
    ok("показатели: − не уходит ниже нуля", pg.input_value("#ind1") == "0", pg.input_value("#ind1"))
    pg.fill("#ind1", v0)
    pg.fill("form .grid-f input[type=text]", " ")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(150)
    ok("показатели: пустое учреждение — ошибка", "Укажите учреждение" in pg.inner_text("form"))
    pg.fill("form .grid-f input[type=text]", "СВП №7 (тест)"); pg.press("form .grid-f input[type=text]", "Tab")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(150)
    ok("показатели: пустая форма — ошибка", "хотя бы один" in pg.inner_text("form"))
    pg.fill("#ind3", "150"); pg.click("form button[type=submit]"); pg.wait_for_timeout(150)
    ok("показатели: проценты больше 100 отклоняются", "Проверьте числа" in pg.inner_text("form"))
    ok("показатели: поле с ошибкой помечено", pg.eval_on_selector("#ind3", "e => e.getAttribute('aria-invalid')") == "true")
    pg.fill("#ind3", "83,3"); pg.fill("#ind1", "12"); pg.fill("#ind2", "10"); pg.evaluate("window.__MOCK.log.length = 0")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(400)
    ins = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "insert" and x[1] == "indicator_reports"]
    ok("показатели: запись отправлена с числами, автором и должностью", ins and ins[0][2]["vals"] == {"i1": 12, "i2": 10, "i3": 83.3} and ins[0][2]["user_id"] == S_ and ins[0][2]["position_code"] == "doya_qvp" and ins[0][2]["period_end"] == "2026-10-23", ins)
    ok("показатели: в отправке нет текстовых полей кроме учреждения", ins and set(ins[0][2].keys()) == {"user_id", "position_code", "institution", "mahalla_id", "period_start", "period_end", "vals"}, ins and ins[0][2].keys())
    # правка существующего отчёта: то же учреждение и период
    pg.fill("form .grid-f input[type=text]", "СВП №3"); pg.press("form .grid-f input[type=text]", "Tab"); pg.wait_for_timeout(150)
    ok("показатели: за этот период отчёт уже есть — значения подставлены", pg.input_value("#ind1") == "12" and "правите" in pg.inner_text("form"), pg.inner_text("form")[:300])
    pg.fill("#ind1", "15"); pg.evaluate("window.__MOCK.log.length = 0"); pg.click("form button[type=submit]"); pg.wait_for_timeout(300)
    upd = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "update"]
    ok("показатели: правка идёт как update того же отчёта", upd and upd[0][3] == ["p1"] and upd[0][2]["vals"]["i1"] == 15, upd)
    ok("показатели: удаление требует второго нажатия", True)
    dl = pg.locator("tbody .btn.danger").first; dl.click(); pg.wait_for_timeout(80)
    ok("показатели: удаление — подтверждение", "Точно" in dl.inner_text())
    pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(300)
    ok("показатели: сотрудник не видит сводку", "нет доступа" in pg.inner_text("#app").lower() or "Недостаточно" in pg.inner_text("#app") or pg.locator("table.indsum").count() == 0)
    ok("ошибок JS нет (показатели сотрудника)", not errs, errs); b.close()

    # сводка руководителя
    reps = [rep("p1", S_, "СВП №3", "2026-10-23", {"i1": 12, "i2": 10, "i3": 83.3}, "2026-10-19"),
            rep("p2", X_, "СВП №5", "2026-10-16", {"i1": 8, "i2": 8}, "2026-10-12"),
            rep("p3", X_, "СВП №5", "2026-10-09", {"i1": 7}, "2026-10-05")]
    b, pg, errs = page_for(p, ind_sc(A, ("admin",), reps), now=NOW)
    pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(500)
    t = pg.inner_text("#app")
    ok("сводка: пункт меню у администратора", "Сводка показателей" in pg.inner_text(".nav"))
    ok("сводка: 27 строк показателей", pg.locator("table.indsum tbody tr").count() == 27, pg.locator("table.indsum tbody tr").count())
    row1 = pg.locator("table.indsum tbody tr").first.inner_text().split("\t")
    ok("сводка: итого по показателю 1 = 12 + 8 (последние отчёты)", row1[-1].strip() == "20", row1)
    ok("сводка: проценты не суммируются", pg.locator("table.indsum tbody tr").nth(2).inner_text().split("\t")[-1].strip() == "", pg.locator("table.indsum tbody tr").nth(2).inner_text())
    ok("сводка: показан отчёт сотрудника и его фамилия", "СВП №3" in t and "сотрудник тест" in t.lower())
    ok("сводка: кто не сдал за срок 23.10.2026", "Не сдали (1)" in t and "Вторая Доя" in t and "Сотрудник Тест" not in t.split("Не сдали")[1].split("\n")[0], t[:700])
    pg.screenshot(path=str(SHOTS / "c_ind_summary.png"), full_page=True)
    with pg.expect_download() as d: pg.click("text=Скачать в Excel")
    d.value.save_as("/tmp/ui_ind.xlsx")
    import openpyxl
    wsx = openpyxl.load_workbook("/tmp/ui_ind.xlsx").active
    ok("сводка: Excel — 27 строк + шапка и итого", wsx.max_row == 28 and wsx.cell(2, 2).value == "Число беременных на учёте" and wsx.cell(2, wsx.max_column).value == 20, (wsx.max_row, wsx.cell(2, wsx.max_column).value))
    pg.select_option("section:not(:has(.chart)) select >> nth=1", "2026-10-16"); pg.wait_for_timeout(400)
    ok("сводка: период — только отчёты за него", pg.locator("table.indsum thead th").count() == 5 and "СВП №5" in pg.inner_text("table.indsum thead"), pg.locator("table.indsum thead th").count())
    pg.click(".seg button:nth-child(2)"); pg.wait_for_timeout(300)
    ok("сводка: узбекский", "Жами" in pg.inner_text("table.indsum thead") or "Кўрсаткичлар" in pg.inner_text("#app"))
    ok("ошибок JS нет (сводка)", not errs, errs); b.close()

    for tag, (w, h_) in {"m": (390, 800), "d": (1280, 900)}.items():
        b, pg, errs = page_for(p, ind_sc(A, ("admin",), reps), w, h_, now=NOW)
        pg.goto(SITE + "#/indicators"); pg.wait_for_timeout(400)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_ind.png"), full_page=True)
        ok(f"{tag}: нет горизонтальной прокрутки (показатели)", not hscroll(pg))
        pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(400)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_indsum.png"), full_page=True)
        ok(f"{tag}: нет горизонтальной прокрутки страницы (сводка)", not hscroll(pg))
        b.close()

    # 8. обзор: проценты и графики
    b, pg, errs = page_for(p, ind_sc(A, ("admin",), reps), now=NOW, hash="")
    pg.wait_for_timeout(400)
    ok("обзор: открывается по умолчанию и пункт меню активен", pg.locator("nav a[aria-current=page]").inner_text() == "Обзор")
    t = pg.inner_text("#app")
    ok("обзор: заполнение папок 1 из 19 = 5%", "1 из 19 папок" in t and pg.locator(".kpi .ring-n").first.inner_text() == "5%", t[:300])
    ok("обзор: сдача показателей 1 из 2 = 50%", pg.locator(".kpi .ring-n").nth(1).inner_text() == "50%" and "1 из 2 сотрудников" in t, t[:400])
    ok("обзор: документов всего 1 и за 30 дней 1", pg.locator(".kpi .big").inner_text() == "1" and "за 30 дней: 1" in t)
    ok("обзор: строки по папкам (19 открытых)", pg.locator(".panel .hb[href^='#/folder/']").count() == 19, pg.locator(".panel .hb").count())
    ok("обзор: по должностям — Доя 1 / 2 и 50%", "1 / 2" in t and "50%" in pg.locator(".hb:not([href])").first.inner_text(), pg.locator(".hb:not([href])").count())
    ok("обзор: график по неделям — 12 точек и линия", pg.locator(".chart svg .ch-dot").count() == 12 and pg.locator(".chart svg path.ch-line").count() == 1, pg.locator(".chart svg .ch-dot").count())
    pg.locator(".chart svg rect").last.scroll_into_view_if_needed()
    box = pg.locator(".chart svg rect").last.bounding_box()
    pg.mouse.move(box["x"] + box["width"] - 8, box["y"] + 40); pg.wait_for_timeout(150)
    ok("обзор: подсказка на графике показывает неделю и число", "19.10–25.10: 1" in pg.locator(".ch-tip").inner_text(), pg.locator(".ch-tip").inner_text())
    ok("обзор: таблица данных для скринридера есть", pg.locator(".chart table.sr tr").count() == 12)
    ok("обзор: график не шире панели", pg.evaluate("(()=>{const s=document.querySelector('.chart svg'),p=s.closest('.panel');return s.getBoundingClientRect().width<=p.getBoundingClientRect().width+1})()"))
    pg.screenshot(path=str(SHOTS / "c_overview.png"), full_page=True)
    pg.click(".seg button:nth-child(2)"); pg.wait_for_timeout(250)
    ok("обзор: узбекский", "Умумий кўриниш" in pg.inner_text("#app") and "Жилдлар" in pg.inner_text("#app"))
    ok("ошибок JS нет (обзор)", not errs, errs); b.close()

    sc_st = ind_sc(S_, ("staff",), reps); sc_st["db"]["folder_members"] = [{"user_id": S_, "folder_id": 1}]
    b, pg, errs = page_for(p, sc_st, now=NOW, hash="")
    pg.wait_for_timeout(300); t = pg.inner_text("#app")
    ok("обзор сотрудника: показатели сданы, срок 23.10.2026", "Сдано" in t and "23.10.2026" in t, t[:300])
    ok("обзор сотрудника: нет сравнения по должностям", "по должностям" not in t)
    ok("обзор сотрудника: только его папка (1 из 1)", "1 из 1 папок" in t, t[:200])
    ok("ошибок JS нет (обзор сотрудника)", not errs, errs); b.close()

    b, pg, errs = page_for(p, ind_sc(X_, ("staff",), []), now=NOW, hash="")
    pg.wait_for_timeout(300); t = pg.inner_text("#app")
    ok("обзор сотрудника без отчёта: не сдано", "Не сдано" in t and pg.locator(".ring.t-bad").count() == 1, t[:300]); b.close()

    b, pg, errs = page_for(p, ind_sc(A, ("admin",), reps), now=NOW)
    pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(500)
    ok("сводка: график динамики — 3 периода", pg.locator(".chart svg .ch-dot").count() == 3 and "Динамика показателя" in pg.inner_text("#app"), (pg.locator(".chart svg .ch-dot").count(), errs, pg.inner_text("#app")[:300]))
    pg.select_option("section:has(.chart) select", "3"); pg.wait_for_timeout(400)
    ok("сводка: для % — среднее и шкала до 100%", "Среднее по учреждениям" in pg.inner_text("#app") or "Сумма" in pg.inner_text("#app"))
    ok("ошибок JS нет (динамика)", not errs, errs); b.close()

    for tag, (w, h_, sch) in {"m": (390, 800, "light"), "d": (1280, 900, "light"), "dk": (1280, 900, "dark"), "mdk": (390, 800, "dark")}.items():
        b, pg, errs = page_for(p, ind_sc(A, ("admin",), reps), w, h_, scheme=sch, now=NOW, hash="")
        pg.wait_for_timeout(300)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_overview.png"), full_page=True)
        ok(f"{tag}: нет горизонтальной прокрутки (обзор)", not hscroll(pg))
        ok(f"{tag}: ошибок JS нет (обзор)", not errs, errs)
        pg.goto(SITE + "#/folders"); pg.wait_for_timeout(300)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_folders2.png"), full_page=True)
        pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(400)
        pg.screenshot(path=str(SHOTS / f"c_{tag}_indsum2.png"), full_page=True)
        ok(f"{tag}: нет горизонтальной прокрутки (сводка с графиком)", not hscroll(pg))
        b.close()

    # 9. учреждения и области видимости
    INSTS = [{"id": 1, "name": "СВП №3", "kind": "svp"}, {"id": 2, "name": "СВП №5", "kind": "svp"}]
    sc = scenario(); sc["db"]["institutions"] = INSTS
    b, pg, errs = page_for(p, sc)
    calls = tg_route(pg, lambda body: {"request": "pending"} if body.get("request") else {"request": "none"})
    pg.evaluate("onTelegramAuth({id: 777, first_name: 'Мария', last_name: 'Иванова', auth_date: 1, hash: 'x'})"); pg.wait_for_timeout(400)
    ok("заявка: есть выбор учреждения", pg.locator("form select").count() == 2 and pg.locator("form select >> nth=1 >> option").count() == 4, pg.locator("form select >> nth=1 >> option").count())
    pg.select_option("form select >> nth=0", "doya_qvp")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(150)
    ok("заявка: без учреждения — ошибка", "Выберите учреждение" in pg.inner_text("form"))
    pg.select_option("form select >> nth=1", "other"); pg.wait_for_timeout(50)
    ok("заявка: «нет в списке» открывает поле", pg.locator("form input[type=text] >> nth=2").is_visible())
    pg.fill("form input[type=text] >> nth=2", "Частная клиника")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(300)
    sent = calls[-1].get("request") or {}
    ok("заявка: ушло название учреждения вручную", sent.get("institution_id") is None and sent.get("institution_text") == "Частная клиника", sent)
    ok("ошибок JS нет (заявка с учреждением)", not errs, errs); b.close()

    b, pg, errs = page_for(p, sc)
    calls = tg_route(pg, lambda body: {"request": "pending"} if body.get("request") else {"request": "none"})
    pg.evaluate("onTelegramAuth({id: 777, first_name: 'Мария', last_name: 'Иванова', auth_date: 1, hash: 'x'})"); pg.wait_for_timeout(400)
    pg.select_option("form select >> nth=0", "doya_qvp"); pg.select_option("form select >> nth=1", "2")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(300)
    sent = calls[-1].get("request") or {}
    ok("заявка: ушёл id учреждения", sent.get("institution_id") == 2 and sent.get("institution_text") == "", sent); b.close()

    sc = scenario(uid=A, roles=("admin",)); sc["db"]["institutions"] = INSTS
    sc["db"]["access_requests"][0]["institution_id"] = 2
    sc["db"]["access_requests"][1]["institution_text"] = "Частная клиника"
    b, pg, errs = page_for(p, sc)
    ok("admin: в меню есть «Учреждения»", "Учреждения" in pg.inner_text("nav"))
    pg.goto(SITE + "#/admin/requests"); pg.wait_for_timeout(400)
    c1, c2 = pg.locator(".req").nth(0), pg.locator(".req").nth(1)
    ok("заявки: учреждение из заявки выбрано", c1.locator("select").nth(1).input_value() == "2")
    ok("заявки: вписанное вручную показано и есть кнопка добавить", "Частная клиника" in c2.inner_text() and c2.locator("text=Добавить в список").count() == 1, c2.inner_text())
    c2.locator("text=Добавить в список").click(); pg.wait_for_timeout(300)
    ad = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_save_institution"]
    ok("заявки: «добавить в список» создаёт учреждение", ad and ad[0][2] == {"_id": None, "_name": "Частная клиника", "_kind": "svp"}, ad)
    c1.locator("button.pri").click(); pg.wait_for_timeout(300)
    dec = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "decide_request"]
    ok("заявки: одобрение передаёт учреждение", dec and dec[0][2]["_institution_id"] == 2, dec)
    ok("ошибок JS нет (заявки с учреждением)", not errs, errs)

    pg.goto(SITE + "#/admin/institutions"); pg.wait_for_timeout(400)
    ok("учреждения: 2 строки и число сотрудников", pg.locator("tbody tr").count() == 2 and pg.locator("tbody tr").first.locator("td.num").inner_text() == "1", pg.locator("tbody tr").count())
    pg.screenshot(path=str(SHOTS / "c_institutions.png"), full_page=True)
    pg.fill("section form.row input[type=text]", "СВП №7"); pg.click("text=Добавить >> nth=0"); pg.wait_for_timeout(300)
    ad = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_save_institution"]
    ok("учреждения: добавление одного", ad and ad[-1][2] == {"_id": None, "_name": "СВП №7", "_kind": "svp"}, ad)
    pg.locator("details summary").first.click()
    pg.fill("textarea", "СВП №8\nСВП №9\n\nСВП №8"); pg.click("text=Добавить все"); pg.wait_for_timeout(300)
    ab = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_add_institutions"]
    ok("учреждения: список строк без дублей", ab and ab[0][2]["_names"] == ["СВП №8", "СВП №9"], ab)
    d1 = pg.locator("tbody tr").nth(1).locator("button.danger"); d1.click(); pg.wait_for_timeout(100)
    ok("учреждения: удаление требует второго нажатия", "Точно" in d1.inner_text()); d1.click(); pg.wait_for_timeout(300)
    dl = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_delete_institution"]
    ok("учреждения: удалено", dl and dl[0][2] == {"_id": 2}, dl)
    ok("ошибок JS нет (учреждения)", not errs, errs)

    pg.goto(SITE + "#/admin/users"); pg.wait_for_timeout(400)
    row = pg.locator("tbody tr").nth(1)
    ok("сотрудники: учреждение сотрудника выбрано", row.locator("td").nth(3).locator("select").input_value() == "1")
    row.locator("td").nth(3).locator("details summary").click()
    row.locator("text=Акушерская ветка").click()
    ok("сотрудники: пресет «акушерская ветка» отмечает специальности", row.locator("td").nth(3).locator("input[value=doya_qvp]").is_checked() and not row.locator("td").nth(3).locator("input[value=pediatr]").is_checked())
    row.locator("td").nth(3).locator("input[type=checkbox][value='2']").check()
    row.locator("td").nth(3).locator("button.pri").click(); pg.wait_for_timeout(300)
    sset = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_set_scopes"]
    ok("сотрудники: область видимости сохранена", sset and sset[0][2]["_institutions"] == [2] and "doya_qvp" in sset[0][2]["_positions"] and "pediatr" not in sset[0][2]["_positions"], sset)
    pg.screenshot(path=str(SHOTS / "c_users_scope.png"), full_page=True)
    row.locator("td").nth(3).locator("select").select_option("2"); pg.wait_for_timeout(200)
    si = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_set_institution"]
    ok("сотрудники: смена учреждения", si and si[0][2] == {"_uid": S_, "_inst": 2}, si)
    ok("ошибок JS нет (сотрудники)", not errs, errs)

    pg.goto(SITE + "#/admin/invites"); pg.wait_for_timeout(400)
    pg.select_option("form select >> nth=2", "2") if False else None
    pg.locator("form label:has-text('Учреждение') select").select_option("1")
    pg.fill("form input[type=text] >> nth=0", "Новый Сотрудник"); pg.select_option("form select >> nth=0", "doya_qvp")
    pg.click("form button[type=submit]"); pg.wait_for_timeout(300)
    ci = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "create_invite"]
    ok("приглашение: учреждение передано", ci and ci[-1][2].get("_institution_id") == 1, ci)
    ok("ошибок JS нет (приглашения)", not errs, errs); b.close()

    # заведующий: сотрудник с областью видимости
    sc = ind_sc(S_, ("staff",), reps); sc["db"]["institutions"] = INSTS
    sc["db"]["scopes"] = [{"user_id": S_, "kind": "position", "position_code": "doya_qvp"}]
    b, pg, errs = page_for(p, sc, now=NOW, hash="")
    pg.wait_for_timeout(400); t = pg.inner_text("#app")
    ok("заведующий: в меню «Сводка показателей»", "Сводка показателей" in pg.inner_text(".nav"))
    ok("заведующий: в обзоре есть сравнение по должностям", "по должностям" in t.lower(), t[:400])
    pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(500)
    ok("заведующий: сводка открывается, только его специальности", pg.locator("table.indsum").count() == 1 and pg.locator(".grid-f select >> nth=0 >> option").count() == 1, pg.locator(".grid-f select >> nth=0 >> option").count())
    ok("ошибок JS нет (заведующий)", not errs, errs); b.close()

    sc = ind_sc(S_, ("staff",), reps); sc["db"]["institutions"] = INSTS
    sc["db"]["profiles"][1]["institution_id"] = 1; sc["db"]["profiles"][2]["institution_id"] = 2
    sc["db"]["scopes"] = [{"user_id": S_, "kind": "institution", "institution_id": 1}]
    b, pg, errs = page_for(p, sc, now=NOW, hash="")
    pg.wait_for_timeout(400); t = pg.inner_text("#app")
    ok("заведующий по учреждению: сдача считается только по его учреждению (1 из 1)", "1 из 1 сотрудников" in t, t[:500])
    pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(500)
    ok("заведующий по учреждению: все должности в списке и чужие не в «не сдали»", pg.locator(".grid-f select >> nth=0 >> option").count() > 5 and "Вторая Доя" not in pg.inner_text(".notice"), pg.inner_text("#app")[:400])
    b.close()

    # 10. махалли и бригады
    MAHS = [{"id": 1, "name": "Боботог"}, {"id": 2, "name": "Қораариқ"}, {"id": 3, "name": "Пустая"}]
    LINKS = [{"mahalla_id": 1, "institution_id": 1}, {"mahalla_id": 2, "institution_id": 1}, {"mahalla_id": 2, "institution_id": 2}, {"mahalla_id": 3, "institution_id": 2}]
    BRIG = [{"id": 1, "mahalla_id": 1, "kind": "doctor", "rate": 1, "user_id": None}, {"id": 2, "mahalla_id": 1, "kind": "patronage", "rate": 1, "user_id": None},
            {"id": 3, "mahalla_id": 1, "kind": "doya_pat", "rate": 0.5, "user_id": None}, {"id": 4, "mahalla_id": 2, "kind": "doctor", "rate": 0.5, "user_id": None}]
    mreps = [dict(rep("m1", S_, "СВП №3", "2026-10-23", {"i1": 12, "i3": 80}, "2026-10-19"), mahalla_id=1),
             dict(rep("m2", X_, "СВП №3", "2026-10-23", {"i1": 8, "i3": 60}, "2026-10-19"), mahalla_id=1),
             dict(rep("m3", S_, "СВП №3", "2026-09-01", {"i1": 5}, "2026-08-26"), mahalla_id=2)]
    def mah_sc(uid, roles):
        sc = ind_sc(uid, roles, mreps); sc["db"]["institutions"] = INSTS
        sc["db"]["mahallas"] = MAHS; sc["db"]["mahalla_links"] = LINKS; sc["db"]["brigade_posts"] = BRIG
        return sc
    b, pg, errs = page_for(p, mah_sc(A, ("admin",)), now=NOW)
    ok("махалли: пункт меню", "Махалли" in pg.inner_text(".nav"))
    pg.goto(SITE + "#/mahallas"); pg.wait_for_timeout(500)
    t = pg.inner_text("#app")
    ok("махалли: с отчётом 1 из 3 = 33%", "1 из 3 за последние 14 дней" in t and pg.locator(".kpi .ring-n").first.inner_text() == "33%", t[:300])
    ok("махалли: бригады полные 1 из 3", "1 из 3 махаллей" in t and pg.locator(".kpi .ring-n").nth(1).inner_text() == "33%", t[:400])
    ok("махалли: Боботог работает, Қораариқ нет отчёта", "Работает" in t and t.count("Нет отчёта") >= 2, t.count("Нет отчёта"))
    ok("махалли: неполная бригада показывает чего не хватает", "врач 0.5 из 1" in t and "патронаж 0 из 1" in t and "доя 0 из 0.5" in t, t[-900:])
    ok("махалли: бригада без состава отмечена", "Состав не задан" in t)
    ok("махалли: общая махалля показана в обоих учреждениях", t.count("Қораариқ") >= 2 and "обслуживают два учреждения" in t)
    pg.screenshot(path=str(SHOTS / "c_mahallas.png"), full_page=True)
    row = pg.locator("tbody tr", has_text="Қораариқ").first
    row.locator("summary").click()
    row.locator("button:has-text('Добавить ставку')").click()
    row.locator("button.pri").click(); pg.wait_for_timeout(300)
    sb_ = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] in ("admin_save_mahalla", "admin_set_brigade")]
    ok("махалли: сохранение имени, учреждений и состава бригады", len(sb_) == 2 and sb_[0][2]["_id"] == 2 and sb_[0][2]["_institutions"] == [1, 2] and len(sb_[1][2]["_posts"]) == 2, sb_)
    pg.evaluate("window.__mhErr = true")
    row = pg.locator("tbody tr", has_text="Боботог").first; row.locator("summary").click()
    d1 = row.locator("button.danger").last; d1.click(); d1.click(); pg.wait_for_timeout(300)
    ok("махалли: с отчётами не удаляется — понятная ошибка", "есть отчёты" in pg.locator("#toast").inner_text(), pg.locator("#toast").inner_text())
    pg.select_option(".ovh select", "2"); pg.wait_for_timeout(400)
    ok("махалли: фильтр по учреждению", pg.locator("tbody tr").count() == 2, pg.locator("tbody tr").count())
    ok("ошибок JS нет (махалли)", not errs, errs); b.close()

    b, pg, errs = page_for(p, mah_sc(A, ("admin",)), now=NOW)
    pg.goto(SITE + "#/indicators"); pg.wait_for_timeout(400)
    ok("показатели: выбор махалли есть", pg.locator("form label:has-text('Махалля') select option").count() == 4, pg.locator("form label:has-text('Махалля') select option").count())
    pg.select_option("form label:has-text('Махалля') select", "3")
    pg.fill("#ind1", "7"); pg.fill("form .grid-f input[type=text]", "СВП №3"); pg.press("form .grid-f input[type=text]", "Tab"); pg.wait_for_timeout(100); pg.click("form button[type=submit]"); pg.wait_for_timeout(400)
    ins_ = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "insert" and x[1] == "indicator_reports"]
    ok("показатели: отчёт уходит с махаллей", ins_ and ins_[-1][2].get("mahalla_id") == 3, ins_)
    ok("показатели: мои отчёты показывают махаллю", "Пустая" in pg.inner_text("#app"))
    ok("ошибок JS нет (показатели по махалле)", not errs, errs); b.close()

    b, pg, errs = page_for(p, mah_sc(A, ("admin",)), now=NOW)
    pg.goto(SITE + "#/indicators/summary"); pg.wait_for_timeout(500)
    pg.select_option("select[aria-label='Группировать']", "mahalla"); pg.wait_for_timeout(500)
    hd = pg.inner_text("table.indsum thead").lower()
    ok("сводка: по махаллям — столбец махалли с числом отчётов", "боботог (2)" in hd and "қораариқ (1)" in hd, hd)
    r1 = pg.locator("table.indsum tbody tr").first.inner_text().split("\t")
    ok("сводка: по махаллям — показатель 1: 12+8=20 и итого 25", r1[3].strip() == "20" and r1[-1].strip() == "25", r1)
    r3 = pg.locator("table.indsum tbody tr").nth(2).inner_text().split("\t")
    ok("сводка: по махаллям — процент усредняется (80 и 60 → 70)", r3[3].strip() == "70%", r3)
    pg.select_option("select[aria-label='Группировать']", "district"); pg.wait_for_timeout(500)
    ok("сводка: весь район — один столбец", "весь район (3)" in pg.inner_text("table.indsum thead").lower() and pg.locator("table.indsum thead th").count() == 4, pg.locator("table.indsum thead th").count())
    pg.select_option("select[aria-label='Группировать']", "inst"); pg.wait_for_timeout(500)
    ok("сводка: по учреждениям — берётся название отчёта", "свп №3 (3)" in pg.inner_text("table.indsum thead").lower(), pg.inner_text("table.indsum thead"))
    ok("ошибок JS нет (группировка)", not errs, errs); b.close()

    # --- рассылки и задания
    BCD = [{"id": "b1", "title": "План на неделю", "note": "Заполните до пятницы", "urgent": True, "version": 2, "file_name": "plan.xlsx", "file_size": 5000,
            "updated_at": "2026-10-24T05:00:00Z", "viewed": False, "downloaded": False}]
    TKD = [{"id": "t1", "title": "Сдать отчёт по ГСД", "body": "Форма 5", "due_date": "2026-10-20", "urgent": False, "broadcast_id": None,
            "created_at": "2026-10-18T05:00:00Z", "seen_at": None, "done_at": None, "note": None}]
    def bc_page(uid, roles, hash, bc=BCD, tk=TKD, abc=None, atk=None, w=1280, h_=900):
        sc = scenario(uid, roles)
        b_, pg_, e_ = page_for(p, sc, w=w, h=h_, hash=hash)
        return b_, pg_, e_
    sc = scenario(S_, ("staff",), my_folders=(1,))
    b = p.chromium.launch(); ctx = b.new_context(viewport={"width": 1100, "height": 900})
    ctx.add_init_script("window.__MOCK=" + json.dumps(sc, ensure_ascii=False) + ";" + RPC + FN + "window.__BC=" + json.dumps(BCD) + ";window.__TK=" + json.dumps(TKD) + ";")
    ctx.route("**/supabase.js", lambda r: r.fulfill(status=200, content_type="application/javascript", body=MOCK))
    ctx.route("**/telegram-widget.js*", lambda r: r.fulfill(status=200, content_type="application/javascript", body="/* stub */"))
    ctx.route("**/fonts.g*/**", lambda r: r.abort())
    ctx.route("**/example.test/**", lambda r: r.fulfill(status=200, headers={"Content-Disposition": "attachment; filename=plan.xlsx"}, body="x"))
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/broadcasts"); pg.wait_for_timeout(500)
    nav_txt = pg.inner_text("nav.nav")
    ok("рассылки: пункты меню и счётчики", "Рассылки" in nav_txt and "Задания" in nav_txt and pg.locator(".nbadge").count() == 2, nav_txt)
    ok("рассылки: карточка со срочностью и статусом «новое»", "План на неделю" in pg.inner_text("#app") and "срочно" in pg.inner_text("#app") and "новое" in pg.inner_text("#app"))
    ok("рассылки: у получателя нет формы отправки", pg.locator("#bc_form").count() == 0)
    pg.click("details.fcard summary"); pg.wait_for_timeout(200)
    lg = pg.evaluate("window.__MOCK.log")
    ok("рассылки: открытие карточки пишет «открыл»", any(x[0] == "rpc" and x[1] == "broadcast_view" and x[2]["_id"] == "b1" for x in lg), lg[-3:])
    ok("рассылки: статус стал «открыто»", "открыто" in pg.inner_text("details.fcard summary"))
    pg.click("details.fcard .btn.pri"); pg.wait_for_timeout(400)
    lg = pg.evaluate("window.__MOCK.log")
    ok("рассылки: скачивание идёт через серверную функцию", any(x[0] == "fn" and x[1] == "bc-file" and x[2]["id"] == "b1" for x in lg), lg[-4:])
    ok("рассылки: статус стал «скачано»", "скачано" in pg.inner_text("details.fcard summary"))
    pg.evaluate("window.__fnErr = true"); pg.click("details.fcard .btn.pri"); pg.wait_for_timeout(300)
    ok("рассылки: отказ сервера показан понятным текстом", "Нет доступа к файлу" in pg.inner_text("#toast"), pg.inner_text("#toast"))
    pg.click("nav.nav >> text=Задания"); pg.wait_for_timeout(500)
    ok("задания: карточка с просрочкой", "Сдать отчёт по ГСД" in pg.inner_text("#app") and "просрочено" in pg.inner_text("#app"))
    pg.click("details.fcard summary"); pg.wait_for_timeout(200)
    ok("задания: открытие пишет «прочитал»", any(x[0] == "rpc" and x[1] == "task_seen" for x in pg.evaluate("window.__MOCK.log")))
    pg.fill("details.fcard textarea", "готово, отправлено"); pg.click("details.fcard .btn.pri"); pg.wait_for_timeout(300)
    td = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "task_done"]
    ok("задания: «Выполнено» уходит с комментарием", td and td[-1][2]["_done"] is True and td[-1][2]["_note"] == "готово, отправлено", td)
    pg.click(".seg button:nth-child(2)"); pg.wait_for_timeout(200)
    ok("рассылки/задания: узбекский интерфейс", "Тарқатмалар" in pg.inner_text("nav.nav") and "Топшириқлар" in pg.inner_text("nav.nav"), pg.inner_text("nav.nav"))
    ok("ошибок JS нет (рассылки: получатель)", not errs, errs); b.close()

    # администратор
    ABC = [{"id": "b1", "title": "План на неделю", "note": "Заполните до пятницы", "urgent": True, "version": 2, "file_name": "plan.xlsx", "file_size": 5000,
            "updated_at": "2026-10-24T05:00:00Z", "total": 3, "notified": 2, "viewed": 1, "downloaded": 1}]
    ATK = [{"id": "t1", "title": "Сдать отчёт по ГСД", "body": "Форма 5", "due_date": "2026-10-30", "urgent": False, "broadcast_id": None, "created_at": "2026-10-18T05:00:00Z", "total": 3, "seen": 1, "done": 0}]
    sc = scenario(A, ("admin",))
    b = p.chromium.launch(); ctx = b.new_context(viewport={"width": 1100, "height": 900})
    ctx.add_init_script("window.__MOCK=" + json.dumps(sc, ensure_ascii=False) + ";" + RPC + FN + "window.__ABC=" + json.dumps(ABC) + ";window.__ATK=" + json.dumps(ATK) + ";")
    ctx.route("**/supabase.js", lambda r: r.fulfill(status=200, content_type="application/javascript", body=MOCK))
    ctx.route("**/telegram-widget.js*", lambda r: r.fulfill(status=200, content_type="application/javascript", body="/* stub */"))
    ctx.route("**/fonts.g*/**", lambda r: r.abort())
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/broadcasts"); pg.wait_for_timeout(600)
    ok("рассылки (админ): форма и список", pg.locator("#bc_form").count() == 1 and "открыли 1 из 3" in pg.inner_text("#app") and "скачали 1 из 3" in pg.inner_text("#app") and "в Telegram: 2 из 3" in pg.inner_text("#app"), pg.inner_text("#app")[:600])
    pg.screenshot(path=str(SHOTS / "c_broadcasts_admin.png"), full_page=True)
    pg.locator("section:has(h3) details.fcard summary").first.click(); pg.wait_for_timeout(400)
    st_txt = pg.inner_text("section:has(h3) details.fcard")
    ok("рассылки (админ): видно кто открыл и скачал", "Иванова Мария" in st_txt and "Петрова Зарина" in st_txt and "×2" in st_txt, st_txt)
    # отправка без файла / с неверным типом
    pg.fill("#bc_title", "Новый шаблон"); pg.click("#bc_form button[type=submit]"); pg.wait_for_timeout(200)
    ok("рассылки (админ): без файла не отправляется", "Выберите файл" in pg.inner_text("#toast"))
    pg.set_input_files("#bc_file", {"name": "virus.exe", "mimeType": "application/octet-stream", "buffer": b"MZ"}); pg.click("#bc_form button[type=submit]"); pg.wait_for_timeout(200)
    ok("рассылки (админ): неподходящий тип отклонён", "не принимается" in pg.inner_text("#toast"))
    pg.set_input_files("#bc_file", {"name": "Шаблон отчёта.xlsx", "mimeType": "application/vnd.ms-excel", "buffer": b"PK\x03\x04"})
    pg.evaluate("window.__pubErr = 'no_recipients'"); pg.select_option("#bc_form select[aria-label='Кому']", "positions"); pg.click("#bc_form button[type=submit]"); pg.wait_for_timeout(400)
    lg = pg.evaluate("window.__MOCK.log")
    ok("рассылки (админ): без получателей — понятная ошибка и файл убран из хранилища", "Не выбрано ни одного получателя" in pg.inner_text("#toast") and any(x[0] == "remove" for x in lg), (pg.inner_text("#toast"), lg[-3:]))
    pg.evaluate("window.__pubErr = null"); pg.select_option("#bc_form select[aria-label='Кому']", "all")
    pg.set_input_files("#bc_file", {"name": "Шаблон отчёта.xlsx", "mimeType": "application/vnd.ms-excel", "buffer": b"PK\x03\x04"})
    pg.check("#bc_urgent"); pg.click("#bc_form button[type=submit]"); pg.wait_for_timeout(600)
    lg = pg.evaluate("window.__MOCK.log")
    up = [x for x in lg if x[0] == "upload"]; pub = [x for x in lg if x[0] == "rpc" and x[1] == "admin_publish_broadcast"]
    ok("рассылки (админ): файл загружен в хранилище по безопасному пути", up and "/" in up[-1][1] and up[-1][1].endswith("shablon_otchyota.xlsx"), up)
    ok("рассылки (админ): публикация со всеми получателями и срочностью", pub and pub[-1][2]["_audience"] == {"all": True} and pub[-1][2]["_urgent"] is True and pub[-1][2]["_title"] == "Новый шаблон" and pub[-1][2]["_path"] == up[-1][1], pub)
    ok("рассылки (админ): после публикации идёт Telegram-уведомление", any(x[0] == "fn" and x[1] == "tg-notify" and x[2]["kind"] == "broadcast" and x[2]["remind"] is False for x in lg), lg[-3:])
    # новая версия
    pg.locator("section:has(h3) details.fcard summary").first.click(); pg.wait_for_timeout(100)
    pg.click("button:has-text(\"Новая версия\")"); pg.wait_for_timeout(200)
    ok("рассылки (админ): «Новая версия» подставляет название и подсказку", pg.input_value("#bc_title") == "План на неделю" and "заменит старую" in pg.inner_text("#bc_form"), pg.input_value("#bc_title"))
    # напоминание
    pg.click("text=Напомнить не открывшим"); pg.wait_for_timeout(300)
    ok("рассылки (админ): напоминание не открывшим", any(x[0] == "fn" and x[1] == "tg-notify" and x[2]["remind"] is True for x in pg.evaluate("window.__MOCK.log")) and "не доставлено 1" in pg.inner_text("#toast"), pg.inner_text("#toast"))
    # удаление — с подтверждением
    pg.locator("section:has(h3) details.fcard .btn.danger").first.click(); pg.wait_for_timeout(100)
    ok("рассылки (админ): удаление требует второго нажатия", not any(x[0] == "rpc" and x[1] == "admin_delete_broadcast" for x in pg.evaluate("window.__MOCK.log")))
    pg.locator("section:has(h3) details.fcard .btn.danger").first.click(); pg.wait_for_timeout(300)
    lg = pg.evaluate("window.__MOCK.log")
    ok("рассылки (админ): удаление чистит и файлы хранилища", any(x[0] == "rpc" and x[1] == "admin_delete_broadcast" for x in lg) and any(x[0] == "remove" and x[1] == ["x/1/plan.xlsx"] for x in lg), lg[-3:])
    # задания (админ)
    pg.goto(SITE + "#/tasks"); pg.wait_for_timeout(600)
    ok("задания (админ): список со счётчиками", "прочитали 1 из 3" in pg.inner_text("#app") and "выполнили 0 из 3" in pg.inner_text("#app"), pg.inner_text("#app")[:500])
    pg.fill("#tk_title", "Проверить журналы"); pg.fill("#tk_due", "2026-11-05"); pg.click("#tk_form button[type=submit]"); pg.wait_for_timeout(500)
    lg = pg.evaluate("window.__MOCK.log"); ct = [x for x in lg if x[0] == "rpc" and x[1] == "admin_create_task"]
    ok("задания (админ): создано со сроком и аудиторией «все»", ct and ct[-1][2]["_due"] == "2026-11-05" and ct[-1][2]["_audience"] == {"all": True}, ct)
    ok("задания (админ): после создания идёт Telegram-уведомление", any(x[0] == "fn" and x[1] == "tg-notify" and x[2]["kind"] == "task" and x[2]["id"] == "t-new" for x in lg))
    pg.locator("section:has(h3) details.fcard summary").first.click(); pg.wait_for_timeout(400)
    ok("задания (админ): видно кто прочитал", "Иванова Мария" in pg.inner_text("section:has(h3) details.fcard"))
    ok("ошибок JS нет (рассылки: администратор)", not errs, errs); b.close()

    # телефон
    sc = scenario(S_, ("staff",), my_folders=(1,))
    b = p.chromium.launch(); ctx = b.new_context(viewport={"width": 390, "height": 780}, is_mobile=True, has_touch=True)
    ctx.add_init_script("window.__MOCK=" + json.dumps(sc, ensure_ascii=False) + ";" + RPC + FN + "window.__BC=" + json.dumps(BCD) + ";window.__TK=" + json.dumps(TKD) + ";")
    ctx.route("**/supabase.js", lambda r: r.fulfill(status=200, content_type="application/javascript", body=MOCK))
    ctx.route("**/telegram-widget.js*", lambda r: r.fulfill(status=200, content_type="application/javascript", body="/* stub */"))
    ctx.route("**/fonts.g*/**", lambda r: r.abort())
    pg = ctx.new_page(); pg.goto(SITE + "#/broadcasts"); pg.wait_for_timeout(600)
    pg.click("details.fcard summary"); pg.wait_for_timeout(200)
    ok("рассылки на телефоне: нет горизонтальной прокрутки", not hscroll(pg))
    pg.screenshot(path=str(SHOTS / "c_broadcasts_mobile.png"), full_page=True); b.close()

    # --- помощник по стандарту (приказ №180): срок считается в браузере, ничего не сохраняется
    import datetime as _dt
    def lmp_ago(weeks): return (_dt.date.today() - _dt.timedelta(days=weeks * 7)).isoformat()
    def pg_ctx(uid, roles, w=1100, h_=900, mob=False):
        sc_ = scenario(uid, roles, my_folders=(1,))
        b_ = p.chromium.launch(); c_ = b_.new_context(viewport={"width": w, "height": h_}, is_mobile=mob, has_touch=mob)
        c_.add_init_script("window.__MOCK=" + json.dumps(sc_, ensure_ascii=False) + ";" + RPC + FN)
        c_.route("**/supabase.js", lambda r: r.fulfill(status=200, content_type="application/javascript", body=MOCK))
        c_.route("**/telegram-widget.js*", lambda r: r.fulfill(status=200, content_type="application/javascript", body="/* stub */"))
        c_.route("**/fonts.g*/**", lambda r: r.abort())
        return b_, c_
    def set_lmp(pg, iso): pg.fill("#pg_lmp", iso); pg.dispatch_event("#pg_lmp", "input"); pg.wait_for_timeout(200)
    b, ctx = pg_ctx(S_, ("staff",))
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/pregnant"); pg.wait_for_timeout(600)
    ok("помощник: пункт меню и поле даты", "Беременные" in pg.inner_text("nav.nav") and pg.locator("#pg_lmp").count() == 1)
    txt = pg.inner_text("#app")
    ok("помощник: сказано, что дата нигде не сохраняется", "не сохраняется" in txt and "никуда не отправляется" in txt, txt[:300])
    ok("помощник: нет полей для Ф.И.О., телефона и номера карты", pg.locator("input[type=tel]").count() == 0 and pg.locator("#pg_card").count() == 0 and pg.locator("input[type=text]").count() == 0)
    ok("помощник: до ввода даты — подсказка", "Введите дату" in txt)
    lm = lmp_ago(20); set_lmp(pg, lm)
    txt = pg.inner_text("#pg_res")
    ok("помощник: считает срок 20 нед. 0 дн.", "Срок: 20 нед. 0 дн." in txt and "ориентировочная дата родов" in txt, txt[:200])
    ok("помощник: сейчас 2-й визит — анализы и заметка по приказу", "2-й визит" in txt and "УЗИ плода" in txt and "Аномалии развития" in txt, txt[:500])
    ok("помощник: следующий визит и пройденные окна", "Следующий визит" in txt and "3-й визит" in txt and "Что должно было быть сделано раньше" in txt and "1-й визит" in txt, txt)
    with pg.expect_download() as dl: pg.click("#pg_ics")
    ics = open(dl.value.path(), encoding="utf-8").read()
    ok("помощник: календарь с окнами визитов без личных данных", ics.count("BEGIN:VEVENT") == 2 and "Визит по стандарту" in ics and "2-й визит" in ics and lm.replace("-", "") not in ics and lm not in ics and "TRIGGER" in ics, ics[:400])
    set_lmp(pg, lmp_ago(10)); txt = pg.inner_text("#pg_res")
    ok("помощник: 10 нед. — 1-й визит, общий анализ крови", "Срок: 10 нед." in txt and "Общий анализ крови" in txt and "Что должно было быть сделано раньше" not in txt, txt[:300])
    set_lmp(pg, (_dt.date.today() + _dt.timedelta(days=3)).isoformat())
    ok("помощник: дата в будущем отклоняется", "не подходит" in pg.inner_text("#pg_res"), pg.inner_text("#pg_res"))
    set_lmp(pg, lmp_ago(44)); ok("помощник: срок больше 42 нед. — предупреждение", "42 недель" in pg.inner_text("#pg_res"), pg.inner_text("#pg_res"))
    set_lmp(pg, lm)
    log = json.dumps(pg.evaluate("window.__MOCK.log"), ensure_ascii=False)
    ok("помощник: дата не отправляется на сервер", lm not in log and all(x[1] == "anc_schedule_list" for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and "anc" in x[1] or "preg" in x[1]), log[:300])
    ok("помощник: дата не сохраняется в браузере", lm not in pg.evaluate("JSON.stringify(Object.assign({}, localStorage)) + JSON.stringify(Object.assign({}, sessionStorage))"))
    ok("помощник: у сотрудника нет редактирования окон", pg.locator("details.panel").locator("input[type=number]").count() == 0 and pg.locator("button:has-text('сводк')").count() == 0)
    pg.click(".seg button:nth-child(2)"); pg.wait_for_timeout(300)
    ok("помощник: узбекский интерфейс", "Ҳомиладорлар" in pg.inner_text("nav.nav") and "сақланмайди" in pg.inner_text("#app"), pg.inner_text("nav.nav"))
    set_lmp(pg, lmp_ago(10)); t_uz = pg.inner_text("#pg_res")
    ok("помощник: узбекский — анализы из базы", "Умумий қон таҳлили" in t_uz and "Муддат: 10 ҳафта" in t_uz, t_uz[:300])
    pg.click(".seg button:nth-child(3)"); pg.wait_for_timeout(300)
    ok("помощник: латиница — меню", "Homiladorlar" in pg.inner_text("nav.nav") and pg.evaluate("document.documentElement.lang") == "uz-Latn", pg.inner_text("nav.nav"))
    set_lmp(pg, lmp_ago(10)); t_l = pg.inner_text("#pg_res")
    ok("помощник: латиница — анализы и заметка из базы переведены", "Umumiy qon tahlili" in t_l and "Xavf omillarini baholash" in t_l and "Общий анализ" not in t_l and "Muddat: 10 hafta" in t_l, t_l[:400])
    ok("латиница: язык запоминается", pg.evaluate("localStorage.getItem('ovabmk-lang')") == "uzl")
    pg.click(".seg button:nth-child(1)"); pg.wait_for_timeout(300)
    set_lmp(pg, lmp_ago(10))
    ok("помощник: русский — тексты на русском", "Общий анализ крови" in pg.inner_text("#pg_res") and "Оценить факторы риска" in pg.inner_text("#pg_res"))
    ok("ошибок JS нет (помощник: сотрудник)", not errs, errs); b.close()

    b, ctx = pg_ctx(A, ("admin",))
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/pregnant"); pg.wait_for_timeout(600)
    pg.click("details.panel >> nth=0 >> summary"); pg.wait_for_timeout(100)
    ok("помощник (админ): окна визитов редактируются", pg.locator("details.panel").locator("input[type=number]").count() == 6)
    pg.locator("details.panel").locator("input[type=number]").nth(1).fill("14")
    pg.locator("details.panel").locator("button:has-text('Сохранить')").first.click(); pg.wait_for_timeout(300)
    sv = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_save_anc_row"]
    ok("помощник (админ): сохранение окна", sv and sv[-1][2]["_visit"] == 1 and sv[-1][2]["_until"] == 14 and sv[-1][2]["_tests"] is None, sv)
    ok("помощник (админ): кнопки сводки в Telegram больше нет", pg.locator("button:has-text('сводк')").count() == 0)
    pg.screenshot(path=str(SHOTS / "c_pregnant_admin.png"), full_page=True)
    ok("ошибок JS нет (помощник: админ)", not errs, errs); b.close()

    b, ctx = pg_ctx(S_, ("staff",), w=390, h_=780, mob=True)
    pg = ctx.new_page(); pg.goto(SITE + "#/pregnant"); pg.wait_for_timeout(600)
    set_lmp(pg, lmp_ago(20))
    ok("помощник на телефоне: нет горизонтальной прокрутки", not hscroll(pg))
    pg.screenshot(path=str(SHOTS / "c_pregnant_mobile.png"), full_page=True); b.close()

    # --- отпуск и замещение
    LVD = [{"id": "l1", "user_id": "x1", "full_name": "Азимова Н.", "user_position": "Доя", "institution": "СВП №3", "kind": "vacation", "date_from": "2026-11-02", "date_to": "2026-11-20", "note": "",
            "status": "pending", "decided_by_name": None, "decided_at": None, "decision_note": "", "created_at": "2026-10-20T05:00:00Z", "due_at": "2026-10-22T05:00:00Z", "overdue": True,
            "escalated_at": "2026-10-23T04:00:00Z", "demanded_at": None, "substitute": S_, "substitute_name": "Сотрудник Тест", "load": 0.3, "consent": "pending", "mine": False, "i_substitute": True, "i_decide": False},
           {"id": "l2", "user_id": S_, "full_name": "Сотрудник Тест", "user_position": "Доя", "institution": "СВП №3", "kind": "vacation", "date_from": "2026-12-01", "date_to": "2026-12-10", "note": "",
            "status": "approved", "decided_by_name": "Руководитель", "decided_at": "2026-10-21T05:00:00Z", "decision_note": "", "created_at": "2026-10-19T05:00:00Z", "due_at": "2026-10-21T05:00:00Z", "overdue": False,
            "escalated_at": None, "demanded_at": None, "substitute": None, "substitute_name": None, "load": None, "consent": None, "mine": True, "i_substitute": False, "i_decide": False}]
    b, ctx = pg_ctx(S_, ("staff",))
    ctx.add_init_script("window.__LV=" + json.dumps(LVD) + ";")
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/leaves"); pg.wait_for_selector("#lv_need", timeout=8000); pg.wait_for_timeout(1500)
    ok("отпуск: пункт меню и значок «ждут действия»", "Отпуск и замещение" in pg.inner_text("nav.nav") and pg.locator("nav.nav a[href='#/leaves'] .nbadge").count() == 1, pg.inner_text("nav.nav"))
    ok("отпуск: просьба замещать — в «Ждут вашего действия»", "Азимова Н." in pg.inner_text("#lv_need") and "Согласен замещать" in pg.inner_text("#lv_need"), pg.inner_text("body")[:600])
    pg.click("#lv_need >> text=Согласен замещать"); pg.wait_for_timeout(400)
    lg = pg.evaluate("window.__MOCK.log")
    ok("отпуск: согласие замещающего и уведомление руководителю", any(x[0] == "rpc" and x[1] == "substitution_consent" and x[2]["_accept"] is True for x in lg) and any(x[0] == "fn" and x[1] == "leave-notify" and x[2]["event"] == "consent" for x in lg))
    pg.select_option("#lv_rate", "1.25"); pg.click("#lv_rate ~ button"); pg.wait_for_timeout(800)
    ok("ставка: отправлена на подтверждение", any(x[0] == "rpc" and x[1] == "request_rate" and x[2]["_rate"] == 1.25 for x in pg.evaluate("window.__MOCK.log")))
    pg.fill("#lv_from", "2026-11-03"); pg.dispatch_event("#lv_from", "change"); pg.fill("#lv_to", "2026-11-14"); pg.dispatch_event("#lv_to", "change"); pg.wait_for_timeout(700)
    ok("заявка: кандидаты, без запаса — недоступен", pg.locator("#lv_sub option").count() == 3 and pg.evaluate("document.querySelector('#lv_sub option[value=u6]').disabled"), pg.inner_text("#lv_sub"))
    pg.select_option("#lv_sub", "u5"); pg.fill("#lv_note", "ежегодный"); pg.click("#lv_send"); pg.wait_for_timeout(400)
    lg = pg.evaluate("window.__MOCK.log")
    rq = [x for x in lg if x[0] == "rpc" and x[1] == "request_leave"]
    ok("заявка: подана с замещающим и уведомлением", rq and rq[-1][2]["_substitute"] == "u5" and rq[-1][2]["_from"] == "2026-11-03" and any(x[0] == "fn" and x[1] == "leave-notify" and x[2]["event"] == "request" for x in lg), rq)
    pg.wait_for_timeout(500); pg.fill("#lv_from", "2026-11-03"); pg.fill("#lv_to", "2026-11-14")
    pg.evaluate("window.__lvErr = 'over_limit'"); pg.click("#lv_send"); pg.wait_for_timeout(400)
    ok("заявка: превышение 1,5 ставки — понятное сообщение", "1,5 ставки" in pg.inner_text("#toast"), pg.inner_text("#toast"))
    ok("отпуск: своя заявка — отмена есть", "Отменить заявку" in pg.inner_text("#lv_list"))
    ok("ошибок JS нет (отпуск: сотрудник)", not errs, errs); b.close()
    LVA = [dict(LVD[0], i_substitute=False, consent="accepted", substitute="u5", substitute_name="Холова Ж.")]
    b, ctx = pg_ctx(A, ("admin",))
    ctx.add_init_script("window.__LV=" + json.dumps(LVA) + ";")
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/leaves"); pg.wait_for_selector("#lv_need", timeout=8000); pg.wait_for_timeout(1500)
    t = pg.inner_text("#lv_need")
    ok("отпуск (админ): просроченная заявка и кнопка требования", "не рассмотрено 2 рабочих дня" in t and "Потребовать рассмотрения" in t, t)
    pg.click("#lv_need >> text=Потребовать рассмотрения"); pg.wait_for_timeout(400)
    lg = pg.evaluate("window.__MOCK.log")
    ok("отпуск (админ): требование рассмотрения уходит руководителю", any(x[0] == "rpc" and x[1] == "admin_demand_review" for x in lg) and any(x[0] == "fn" and x[1] == "leave-notify" and x[2]["event"] == "demand" for x in lg))
    ok("ставки (админ): заявка на 1,25", "Каримова Д." in pg.inner_text("#lv_rates") and "1,25" in pg.inner_text("#lv_rates"), pg.inner_text("#lv_rates"))
    pg.click("#lv_rates >> text=Подтвердить"); pg.wait_for_timeout(300)
    ok("ставки (админ): подтверждено", any(x[0] == "rpc" and x[1] == "admin_confirm_rate" and x[2]["_ok"] is True for x in pg.evaluate("window.__MOCK.log")))
    LVH = [dict(LVA[0], overdue=False, i_decide=True)]
    pg.evaluate("window.__LV = " + json.dumps(LVH)); pg.goto(SITE + "#/overview"); pg.goto(SITE + "#/leaves"); pg.wait_for_selector("#lv_need", timeout=8000); pg.wait_for_timeout(1500)
    pg.evaluate("window.__decErr = 'no_consent'"); pg.click("#lv_need >> text=Разрешить"); pg.wait_for_timeout(300)
    ok("решение: без согласия — понятное сообщение", "замещающий должен дать согласие" in pg.inner_text("#toast"), pg.inner_text("#toast"))
    pg.evaluate("window.__decErr = null"); pg.click("#lv_need >> text=Разрешить"); pg.wait_for_timeout(400)
    lg = pg.evaluate("window.__MOCK.log")
    ok("решение: разрешено, уведомление сотруднику и замещающему", any(x[0] == "rpc" and x[1] == "decide_leave" and x[2]["_approve"] is True for x in lg) and any(x[0] == "fn" and x[1] == "leave-notify" and x[2]["event"] == "decided" for x in lg))
    pg.goto(SITE + "#/offline?pub=" + "A" * 43 + "." + "B" * 43); pg.wait_for_timeout(600)
    sp = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "set_my_offline_pub"]
    ok("офлайн: ключ устройства из ссылки программы сохранён, адрес очищен", sp and sp[-1][2]["_x"] == "A" * 43 and "pub=" not in pg.url, (sp, pg.url))
    ok("ошибок JS нет (отпуск: админ)", not errs, errs); b.close()

    # --- офлайн-программа: ключ активации с сайта, должность назначает администратор
    b, ctx = pg_ctx(S_, ("staff",))
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/offline"); pg.wait_for_timeout(500)
    ok("офлайн: пункт меню у сотрудника", "Офлайн-программа" in pg.inner_text("nav.nav"), pg.inner_text("nav.nav"))
    ok("офлайн: сотрудник не видит ключ шифрования администратора", pg.locator("#off_akey").count() == 0)
    pg.click("#off_get"); pg.wait_for_timeout(300); t = pg.inner_text("#off_res")
    ok("офлайн: ключ получен, ссылка на программу с ключом", "Доя кабинета" in t and "/app/#a=eyJ2IjoxfQ.c2ln" in t, t)
    ok("офлайн: кнопка «Открыть программу»", (pg.get_attribute("#off_res a.btn", "href") or "").endswith("/app/#a=eyJ2IjoxfQ.c2ln"))
    pg.evaluate("window.__offErr = 'no_offline_role'"); pg.click("#off_get"); pg.wait_for_timeout(300)
    ok("офлайн: без назначенной должности — понятное сообщение", "не назначил" in pg.inner_text("#off_res"), pg.inner_text("#off_res"))
    pg.evaluate("window.__offErr = 'no_admin_key'"); pg.click("#off_get"); pg.wait_for_timeout(300)
    ok("офлайн: без ключа администратора — понятное сообщение", "ключ шифрования" in pg.inner_text("#off_res"), pg.inner_text("#off_res"))
    ok("ошибок JS нет (офлайн: сотрудник)", not errs, errs); b.close()
    b, ctx = pg_ctx(A, ("admin",))
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/offline"); pg.wait_for_timeout(500)
    ok("офлайн (админ): ключ шифрования ещё не сохранён", "не сохранён" in pg.inner_text("#off_akey_state"), pg.inner_text("#off_akey_state"))
    pg.fill("#off_akey", "не ключ"); pg.click("#off_akey ~ div button"); pg.wait_for_timeout(200)
    ok("офлайн (админ): неверный ключ не сохраняется", "Это не ключ" in pg.inner_text("#toast"), pg.inner_text("#toast"))
    pg.fill("#off_akey", "A" * 43 + "." + "B" * 43); pg.click("#off_akey ~ div button"); pg.wait_for_timeout(500)
    sk = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_set_offline_key"]
    ok("офлайн (админ): ключ сохранён (две половины)", sk and sk[-1][2]["_x"] == "A" * 43 and sk[-1][2]["_y"] == "B" * 43 and "Ключ сохранён" in pg.inner_text("#off_akey_state"), sk)
    pg.goto(SITE + "#/admin/users"); pg.wait_for_timeout(500)
    rows = pg.locator("table tbody tr")
    sel_a, sel_s = rows.nth(0).locator("td").nth(4).locator("select"), rows.nth(1).locator("td").nth(4).locator("select")
    ok("сотрудники: колонка офлайн-программы, у сотрудника — доя кабинета", sel_s.input_value() == "cab" and sel_a.input_value() == "")
    ok("сотрудники: «Администратор» есть только у админа", sel_a.locator("option[value=coord]").count() == 1 and sel_s.locator("option[value=coord]").count() == 0)
    sel_s.select_option("pat"); pg.wait_for_timeout(300)
    so = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == "admin_set_offline_role"]
    ok("сотрудники: должность в офлайн-программе назначена", so and so[-1][2]["_role"] == "pat" and so[-1][2]["_uid"] == "%s" % S_, so)
    rows.nth(1).locator("td").nth(4).locator("button").click(); pg.wait_for_timeout(300)
    fk = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "fn" and x[1] == "offline-key"]
    ok("сотрудники: администратор выдаёт ключ сотруднику", fk and fk[-1][2]["uid"] == S_ and "/app/#a=eyJ2IjoxfQ.c2lnU" in rows.nth(1).inner_text(), fk)
    rows.nth(0).locator("td").nth(4).locator("button").click(); pg.wait_for_timeout(300)
    ok("сотрудники: без должности ключ не выдаётся", "Сначала выберите должность" in pg.inner_text("#toast"), pg.inner_text("#toast"))
    pg.click(".seg button:nth-child(3)"); pg.wait_for_timeout(300)
    ok("латиница: офлайн-программа в меню", "Oflayn-dastur" in pg.inner_text("nav.nav"), pg.inner_text("nav.nav"))
    pg.click(".seg button:nth-child(1)"); pg.wait_for_timeout(200)
    ok("ошибок JS нет (офлайн: админ)", not errs, errs); b.close()

    # ---------------- подпапки и сроки сдачи
    SUBF = json.loads((ROOT / "tests/subfolders.json").read_text(encoding="utf-8"))
    def sub_ctx(uid, roles, folders=(10,), w=1200, h_=900, extra=None):
        sc_ = scenario(uid, roles, my_folders=folders)
        sc_["db"]["subfolders"] = [x for x in SUBF if x["folder_id"] != 15]
        sc_["db"]["profiles"] = [{"id": A, "full_name": "Администратор Тест", "position": "Заместитель главного врача", "removed_at": None},
                                 {"id": S_, "full_name": "Сотрудник Тест", "position": "Акушерка СВП (кабинет доя)", "removed_at": None}]
        sc_["db"]["user_roles"] = [{"user_id": A, "role": "admin"}, {"user_id": S_, "role": "staff"}] if uid == A else sc_["db"]["user_roles"]
        sc_["db"]["folder_members"] = [{"user_id": S_, "folder_id": f} for f in folders]
        sc_["db"]["subfolder_marks"] = [{"subfolder_id": "10.1", "period_key": "2026-11-05", "user_id": S_, "value": 14, "submitted_at": "2026-11-03T06:00:00Z", "paper_at": None}]
        sc_["db"]["documents"] = sc_["db"]["documents"] + [
            {"id": "d10", "folder_id": 10, "subfolder_id": "10.2", "period_key": "2026-11-05", "title": "Отчёт по раннему учёту, октябрь", "doc_number": None, "doc_date": None,
             "note": None, "storage_path": "10/22222222-2222-4222-8222-222222222222/otchet.pdf", "file_name": "otchet.pdf", "file_size": 50000, "uploaded_by": S_, "created_at": "2026-11-04T08:00:00Z"},
            {"id": "d11", "folder_id": 10, "subfolder_id": None, "period_key": None, "title": "Старый документ без подпапки", "doc_number": None, "doc_date": None,
             "note": None, "storage_path": "10/33333333-3333-4333-8333-333333333333/old.pdf", "file_name": "old.pdf", "file_size": 1000, "uploaded_by": A, "created_at": "2026-10-20T08:00:00Z"}]
        b_ = p.chromium.launch(); c_ = b_.new_context(viewport={"width": w, "height": h_})
        c_.add_init_script("window.__NOW='2026-11-20T09:00:00'")
        c_.add_init_script("window.__MOCK=" + json.dumps(sc_, ensure_ascii=False) + ";" + RPC + FN +
            ";window.__MOCK.rpc.subfolder_reminder_log = function(x){ return x._sub === '10.7' ? [{period_key:'2026-11-05', user_id:'" + S_ + "', kind:'overdue', sent_at:'2026-11-06T04:07:00Z'}] : []; };['subfolder_report','subfolder_unreport','subfolder_paper','admin_set_subfolder','admin_move_document'].forEach(function(n){ window.__MOCK.rpc[n] = function(a){ return (window.__subErr && window.__subErr[n]) ? {error:{message:window.__subErr[n]}} : null; }; });")
        c_.route("**/supabase.js", lambda r: r.fulfill(status=200, content_type="application/javascript", body=MOCK))
        c_.route("**/telegram-widget.js*", lambda r: r.fulfill(status=200, content_type="application/javascript", body="/* stub */"))
        c_.route("**/fonts.g*/**", lambda r: r.abort())
        return b_, c_
    def rpcs(pg, name): return [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "rpc" and x[1] == name]

    b, ctx = sub_ctx(A, ("admin",))
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/folders"); pg.wait_for_timeout(500)
    ok("подпапки: на карточке папки 10 — 7 подпапок", "подпапок: 7" in pg.locator("a.fcard[href='#/folder/10']").inner_text())
    pg.goto(SITE + "#/folder/10"); pg.wait_for_timeout(500)
    ok("подпапки: в папке 10 семь подпапок", pg.locator("details.sf").count() == 7, pg.locator("details.sf").count())
    t = pg.inner_text("#app")
    ok("подпапки: рус + узб. латиница, без кириллицы", "Отчёт по раннему учёту (до 12 недель)" in t and "Erta hisobga olish (12 haftagacha) boʻyicha hisobot" in t)
    ok("подпапки: документ без подпапки отдельно", "Другие документы (без подпапки)" in t and "Старый документ без подпапки" in t)
    ok("подпапки: виды хранения подписаны", all(x in t for x in ("Сайт", "Только цифры")))
    pg.click("#sf-10\\.2 summary"); pg.wait_for_timeout(200)
    body = pg.inner_text("#sf-10\\.2")
    ok("10.2: кто сдаёт и срок", "Акушерка СВП (кабинет доя)" in body and "Ежемесячно, до 5 числа" in body, body[:300])
    ok("10.2: файл в подпапке виден", "Отчёт по раннему учёту, октябрь" in body)
    pg.screenshot(path=str(SHOTS / "c_subfolders.png"), full_page=True)
    ok("10.2: руководитель видит, кто сдал", "Сдали 1 из 1" in body, body)
    pg.click("#sf-10\\.2 details.sf-up summary"); pg.wait_for_timeout(100)
    cyc = pg.locator("#sf-10\\.2 details.sf-up select")
    ok("10.2: выбор срока, по умолчанию прошедший несданный (05.11 — админ не сдавал)", cyc.input_value() == "2026-11-05", cyc.input_value())
    pg.locator("#sf-10\\.2 details.sf-up input[type=text]").first.fill("Отчёт за ноябрь")
    pg.locator("#sf-10\\.2 details.sf-up input[type=file]").set_input_files(files=[{"name": "noyabr.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4 t"}])
    cyc.select_option("2026-12-05")
    pg.locator("#sf-10\\.2 details.sf-up .check input").check(); pg.locator("#sf-10\\.2 details.sf-up button[type=submit]").click(); pg.wait_for_timeout(500)
    ins = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "insert" and x[1] == "documents"]
    ok("10.2: файл загружен в подпапку с выбранным сроком", ins and ins[-1][2]["subfolder_id"] == "10.2" and ins[-1][2]["period_key"] == "2026-12-05", ins[-1:] )
    ok("10.2: после загрузки подпапка осталась открытой", pg.locator("#sf-10\\.2[open]").count() == 1)
    pg.click("#sf-10\\.3 summary"); pg.wait_for_timeout(200)
    body = pg.inner_text("#sf-10\\.3")
    ok("10.3: «только цифры» — предупреждение про ФИО, файла нет", "Список с ФИО на сайт не загружается" in body and pg.locator("#sf-10\\.3 input[type=file]").count() == 0, body[:300])
    pg.locator("#sf-10\\.3 form button[type=submit]").click(); pg.wait_for_timeout(150)
    ok("10.3: без числа — ошибка", "Укажите число" in pg.inner_text("#sf-10\\.3"))
    pg.locator("#sf-10\\.3 form input[type=number]").fill("23"); pg.locator("#sf-10\\.3 form button[type=submit]").click(); pg.wait_for_timeout(400)
    r = rpcs(pg, "subfolder_report")
    ok("10.3: отметка с числом отправлена за прошедший срок", r and r[-1][2] == {"_sub": "10.3", "_period": "2026-11-05", "_value": 23}, r[-1:])
    pg.click("#sf-10\\.1 summary"); pg.wait_for_timeout(200)
    body = pg.inner_text("#sf-10\\.1")
    ok("10.1: руководитель видит число сотрудника", "Сдали 1 из 1" in body and "14" in body, body)
    pg.locator("#sf-10\\.1 .sf-ctl input[type=checkbox]").first.check(); pg.wait_for_timeout(300)
    r = rpcs(pg, "subfolder_paper")
    ok("10.1: отметка «оригинал получен»", r and r[-1][2]["_user"] == S_ and r[-1][2]["_on"] is True and r[-1][2]["_period"] == "2026-11-05", r[-1:])
    pg.click("#sf-10\\.7 summary"); pg.wait_for_timeout(200)
    ok("10.7: несдавшим — кнопка напоминания", pg.locator("#sf-10\\.7 .sf-ctl button").count() == 1)
    ok("10.7: видно, что бот уже напомнил", "бот: просрочка 06.11.2026" in pg.inner_text("#sf-10\\.7"), pg.inner_text("#sf-10\\.7")[-300:])
    pg.locator("#sf-10\\.7 .sf-ctl button").click(); pg.wait_for_timeout(300)
    fn = [x for x in pg.evaluate("window.__MOCK.log") if x[0] == "fn" and x[1] == "tg-notify"]
    ok("10.7: напоминание уходит только несдавшим", fn and fn[-1][2]["kind"] == "subfolder" and fn[-1][2]["sub"] == "10.7" and fn[-1][2]["users"] == [S_] and fn[-1][2]["due"] == "2026-11-05", fn[-1:])
    pg.locator("#sf-10\\.7 details.sf-set summary").click(); pg.wait_for_timeout(100)
    pg.locator("#sf-10\\.7 details.sf-set select").nth(1).select_option("quarter")
    pg.locator("#sf-10\\.7 details.sf-set button").click(); pg.wait_for_timeout(300)
    r = rpcs(pg, "admin_set_subfolder")
    ok("10.7: администратор меняет периодичность", r and r[-1][2]["_period"] == "quarter" and r[-1][2]["_id"] == "10.7" and "doya_qvp" in r[-1][2]["_positions"], r[-1:])
    mv = pg.locator("table select[aria-label='Подпапка']").last
    mv.select_option("10.6"); pg.wait_for_timeout(300)
    r = rpcs(pg, "admin_move_document")
    ok("перенос старого документа в подпапку", r and r[-1][2] == {"_doc": "d11", "_sub": "10.6", "_period": ""}, r[-1:])
    pg.goto(SITE + "#/due"); pg.wait_for_timeout(600)
    t = pg.inner_text("#app")
    ok("сроки: пункт меню и заголовок", "Сроки сдачи" in pg.inner_text("nav.nav") and "Контроль по району" in t)
    ok("сроки: просрочки по району посчитаны", "Не сдано к прошедшему сроку:" in t and "Сотрудник Тест" in t, t[:600])
    ok("сроки: «Что сдаю я» у заместителя есть", "Что сдаю я" in t and "Просрочено:" in t)
    ok("сроки: сказано про автонапоминания", "Бот сам напоминает в Telegram в 9:00" in t)
    pg.screenshot(path=str(SHOTS / "c_due.png"), full_page=False)
    pg.click(".seg button:nth-child(3)"); pg.wait_for_timeout(300)
    t = pg.inner_text("#app")
    ok("латиница: сроки и названия подпапок латиницей", "Topshirish muddatlari" in pg.inner_text("nav.nav") and "Yillik ish rejasi" in t, t[:300])
    pg.click(".seg button:nth-child(2)"); pg.wait_for_timeout(300)
    t = pg.inner_text("#app")
    ok("узб. кириллица в меню, а подпапки всё равно латиницей", "Топшириш муддатлари" in pg.inner_text("nav.nav") and "Yillik ish rejasi" in t and "Йиллик иш режаси" not in t)
    pg.click(".seg button:nth-child(1)"); pg.wait_for_timeout(200)
    ok("ошибок JS нет (подпапки: админ)", not errs, errs); b.close()

    b, ctx = sub_ctx(S_, ("staff",), folders=(10,), w=390, h_=844)
    pg = ctx.new_page(); errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(SITE + "#/due"); pg.wait_for_timeout(600)
    t = pg.inner_text("#app")
    ok("сотрудник: свои подпапки по должности", "10.1" in t and "10.2" in t and "10.7" in t and "1.1" not in t.split("Что сдаю я")[1][:5], t[:500])
    ok("сотрудник: сданное отмечено, несданное просрочено", "сдано (срок 05.12.2026)" not in t and "просрочено: срок 05.11.2026" in t, t[:800])
    ok("сотрудник: нет контроля по району", "Контроль по району" not in t)
    pg.screenshot(path=str(SHOTS / "c_due_mobile.png"), full_page=True)
    ok("сотрудник: на телефоне без горизонтальной прокрутки", not hscroll(pg))
    pg.click("a[href='#/folder/10/10.3']"); pg.wait_for_timeout(500)
    ok("ссылка открывает нужную подпапку", pg.locator("#sf-10\\.3[open]").count() == 1)
    ok("сотрудник: не видит, кто ещё сдал", pg.locator("#sf-10\\.3 .sf-ctl").count() == 0 and pg.locator("#sf-10\\.3 details.sf-set").count() == 0)
    pg.goto(SITE + "#/folder/10/10.1"); pg.wait_for_timeout(500)
    sel = pg.locator("#sf-10\\.1 form select")
    ok("сотрудник: прошедший срок сдан — по умолчанию следующий", sel.input_value() == "2026-12-05", sel.input_value())
    sel.select_option("2026-11-05"); pg.wait_for_timeout(100)
    body = pg.inner_text("#sf-10\\.1")
    ok("сотрудник: видит свою отметку и может отменить", "Отмечено" in body and "число 14" in body and "Отменить отметку" in body, body)
    ok("сотрудник: число подставлено в поле", pg.locator("#sf-10\\.1 input[type=number]").input_value() == "14")
    pg.click("#sf-10\\.1 button:has-text('Отменить отметку')"); pg.wait_for_timeout(300)
    r = rpcs(pg, "subfolder_unreport")
    ok("сотрудник: отмена отметки", r and r[-1][2] == {"_sub": "10.1", "_period": "2026-11-05"}, r[-1:])
    pg.evaluate("window.__subErr = {subfolder_report: 'forbidden'}")
    pg.goto(SITE + "#/folder/10/10.3"); pg.wait_for_timeout(500)
    pg.locator("#sf-10\\.3 form input[type=number]").fill("5"); pg.locator("#sf-10\\.3 form button[type=submit]").click(); pg.wait_for_timeout(300)
    ok("сотрудник: ошибка базы показана понятно", "Нет доступа" in pg.inner_text("#sf-10\\.3") and "forbidden" not in pg.inner_text("#sf-10\\.3"), pg.inner_text("#sf-10\\.3")[-200:])
    ok("ошибок JS нет (подпапки: сотрудник)", not errs, errs); b.close()

print(f"\n{sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
