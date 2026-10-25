import os, time, re, sys, json
from playwright.sync_api import sync_playwright
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from tok import init_script, token, PUB
HERE=os.path.dirname(os.path.abspath(__file__)); D=HERE+"/data/"; URL="file://"+os.path.abspath(HERE+"/../dist/OvaBMQ_Offline.html")
PW="Testpass123!"
ok=0; fails=[]
def check(name,cond,extra=""):
    global ok
    if cond: ok+=1
    else: fails.append(name+" "+str(extra)); print("FAIL",name,extra)
def wait(fn,t=60):
    end=time.time()+t
    while time.time()<end:
        if fn(): return
        time.sleep(0.1)
    raise Exception("timeout")
# v2 новее v1
t=time.time()
os.utime(D+"Baza_01_Doya_QVP_SVP3.xlsx",(t-5000,t-5000)); os.utime(D+"Baza_01_Doya_QVP_SVP3_v2.xlsx",(t-100,t-100))
with sync_playwright() as p:
    br=p.chromium.launch(); ctx=br.new_context(viewport={"width":1100,"height":900},accept_downloads=True)
    ctx.add_init_script(init_script("coord"))
    reqs=[]; ctx.on("request",lambda r: reqs.append(r.url))
    pg=ctx.new_page(); errs=[]
    pg.on("pageerror",lambda e: errs.append(str(e))); pg.on("console",lambda m: errs.append(m.text) if m.type=="error" and "Content Security" not in m.text else None)
    pg.goto(URL)
    # --- пароль
    pg.fill("#pwA","short"); pg.fill("#pwB","short"); pg.click("#authGo")
    check("short pw rejected", "не короче" in pg.inner_text("#auth"))
    pg.fill("#pwA",PW); pg.fill("#pwB","other"); pg.click("#authGo")
    check("mismatch pw rejected", "не совпадают" in pg.inner_text("#auth"))
    pg.fill("#pwB",PW); pg.click("#authGo"); pg.wait_for_selector("#app:not(.hidden)")
    check("app opens", pg.is_visible("#nav"))
    check("empty home hint", "С чего начать" in pg.inner_text("#main"))
    # --- импорт
    pg.click("nav >> text=Файлы")
    files=[D+f for f in ["Baza_01_Doya_QVP_SVP3.xlsx","Baza_01_Doya_QVP_SVP5.xlsx","Baza_04_Pediatr_SVP3.xlsx","Baza_02_Akusher_RB.xlsx","Baza_01_Doya_QVP_SVP3_v2.xlsx"]]
    pg.set_input_files("#fileIn",files); wait(lambda: len(pg.query_selector_all(".ln"))==5 and "⏳" not in pg.inner_text(".log"))
    lines=[x.inner_text() for x in pg.query_selector_all(".ln")]
    check("5 files processed", len(lines)==5, lines)
    check("4+1 ok", sum(l.startswith("✔") for l in lines)==5, lines)
    check("doya row count 10 in log", "строк данных: 10" in lines[0], lines[0])
    # дубликаты
    pg.set_input_files("#fileIn",[D+"Baza_04_Pediatr_SVP3.xlsx"]); pg.wait_for_selector(".ln.warn")
    check("duplicate skipped", "уже загружен" in pg.inner_text(".ln.warn"))
    pg.set_input_files("#fileIn",[D+"Baza_01_Doya_QVP_SVP3.xlsx"]); wait(lambda: pg.query_selector(".ln.warn") is not None)
    check("older version skipped", "более новая" in pg.inner_text(".ln") or "уже загружен" in pg.inner_text(".ln"), pg.inner_text(".ln"))
    check("4 files in list (v1 replaced)", len(pg.query_selector_all("#main table tbody tr"))==4, len(pg.query_selector_all("#main table tbody tr")))
    # --- отправители / статусы
    pg.click("nav >> text=Отправители")
    rows=pg.query_selector_all("#main table tbody tr"); check("4 senders", len(rows)==4, len(rows))
    txt=pg.inner_text("#main")
    check("SVP5 late", re.search(r"СВП №5.*?просрочено",txt,re.S) is not None)
    check("RB soon", "скоро срок" in txt)
    check("badge counts", txt.count("в срок")>=2, txt.count("в срок"))
    # --- сводка
    pg.click("nav >> text=Сводка"); t=pg.inner_text("#main")
    check("home late panel", "Не сдали в срок (1)" in t and "СВП №5" in t, t[:400])
    check("home soon panel", "Скоро срок (1)" in t)
    check("upcoming deadlines", "Ближайшие сроки" in t)
    # напоминание в буфер
    ctx.grant_permissions(["clipboard-read","clipboard-write"]) if False else None
    pg.click("text=Напомнить >> nth=0"); check("toast copied", pg.is_visible(".toast"))
    with pg.expect_download() as dl: pg.click("text=Календарь (.ics) для телефона")
    path=dl.value.path(); ics=open(path,encoding="utf-8").read()
    check("ics ok", ics.startswith("BEGIN:VCALENDAR") and "BEGIN:VEVENT" in ics and "VALARM" in ics and "Доя" in ics)
    check("ics has Fri 12:00", "DTSTART:20261030T120000" in ics, ics[:600])
    # --- данные
    pg.click("nav >> text=Данные"); pg.wait_for_selector("#main table")
    opts=pg.eval_on_selector_all("#d_sub option","o=>o.map(x=>x.textContent)")
    check("data options 4", len(opts)==4, opts)
    svp3=[o for o in opts if "Доя" in o and "№3" in o][0]; val=pg.eval_on_selector("#d_sub",f"(s)=>[...s.options].find(o=>o.textContent==={json.dumps(svp3)}).value")
    pg.select_option("#d_sub",val); pg.wait_for_timeout(300)
    check("latest SVP3 has 12 rows", "Записей: 12" in pg.inner_text("#main"), pg.inner_text("#main")[:300])
    pg.fill("#d_flt","Образцова"); pg.wait_for_timeout(200)
    check("filter works", "Записей: 1 (из 12)" in pg.inner_text("#main"))
    with pg.expect_download() as dl2: pg.click("text=Скачать в Excel")
    import openpyxl
    dl2.value.save_as("/tmp/t_sheet.xlsx"); wb=openpyxl.load_workbook("/tmp/t_sheet.xlsx"); check("sheet export", wb.active.max_row==2, wb.active.max_row)
    # --- поиск
    pg.click("nav >> text=Поиск"); pg.fill("#s_q","Образцова"); pg.wait_for_selector("details.res")
    check("search 3 hits", "Найдено записей: 3" in pg.inner_text("#main"), pg.inner_text("#main")[:200])
    pg.fill("#s_q","30000000000003"); pg.wait_for_timeout(500)
    check("search by pinfl", "Найдено записей:" in pg.inner_text("#main") and "Найдено записей: 0" not in pg.inner_text("#main"), pg.inner_text("#main")[:200])
    pg.fill("#s_q","нетттакого"); pg.wait_for_timeout(500); check("search none", "Ничего не найдено" in pg.inner_text("#main"))
    pg.fill("#s_q","Образцова"); pg.wait_for_selector("details.res"); pg.click("details.res >> nth=0 >> summary")
    check("details expanded shows PINFL", re.search(r"\d{14}",pg.inner_text("details.res[open]")) is not None)
    # --- показатели
    pg.click("nav >> text=Показатели"); pg.select_option("#i_role","1"); pg.wait_for_timeout(200)
    t=pg.inner_text("#main"); check("indicators 2 senders", "СВП №3" in t and "СВП №5" in t)
    first=pg.query_selector_all("#main table tbody tr")[0].inner_text().split("\t")
    check("indicator total 18", first[-1].strip()=="18", first)
    # --- статус в Excel
    pg.click("nav >> text=Отправители")
    with pg.expect_download() as dl3: pg.click("text=Статус в Excel")
    dl3.value.save_as("/tmp/t_status.xlsx"); wb=openpyxl.load_workbook("/tmp/t_status.xlsx"); check("status xlsx", wb.active.max_row==5 and wb.active["A1"].value=="Должность", wb.active.max_row)
    # --- добавить вручную
    pg.click("text=Добавить отправителя"); pg.select_option(".modal select","3"); pg.fill(".modal #f_inst","СВП №9 (вручную)"); pg.click(".modal button.pri")
    pg.wait_for_selector("text=СВП №9 (вручную)")
    pg.click("nav >> text=Сводка"); check("manual sender late (since last deadline)", "Не сдали в срок" in pg.inner_text("#main"))
    # --- резервная копия
    pg.click("nav >> text=Настройки")
    with pg.expect_download() as dl4: pg.click("text=Сохранить копию")
    bak=dl4.value.path(); btxt=open(bak,encoding="utf-8").read()
    check("backup is encrypted (no names)", "Образцова" not in btxt and "СВП" not in btxt and "ovabmk-offline" in btxt)
    # --- блокировка и пароль
    pg.click("text=Заблокировать сейчас"); pg.wait_for_selector("#auth:not(.hidden)")
    check("locked hides data", pg.inner_text("#main")=="" and "Образцова" not in pg.content().split("<script>")[0])
    pg.fill("#pwA","wrongwrong"); pg.click("#authGo"); wait(lambda: "Неверный пароль" in pg.inner_text("#auth")); check("wrong pw", True)
    pg.fill("#pwA",PW); pg.click("#authGo"); pg.wait_for_selector("#app:not(.hidden)")
    # --- перезагрузка: сохранение в IndexedDB
    pg.reload(); pg.wait_for_selector("#authGo"); check("vault persists after reload", "Открыть" in pg.inner_text("#authGo"))
    pg.fill("#pwA",PW); pg.click("#authGo"); pg.wait_for_selector("#app:not(.hidden)"); pg.click("nav >> text=Отправители")
    check("senders persist", len(pg.query_selector_all("#main table tbody tr"))==5, len(pg.query_selector_all("#main table tbody tr")))
    # --- смена пароля
    pg.click("nav >> text=Настройки"); pg.click("text=Сменить пароль")
    inputs=pg.query_selector_all(".modal input"); inputs[0].fill(PW); inputs[1].fill("NewPass456!"); inputs[2].fill("NewPass456!"); pg.click(".modal button.pri")
    pg.wait_for_selector("text=Пароль изменён")
    pg.click("text=Заблокировать сейчас"); pg.fill("#pwA",PW); pg.click("#authGo"); wait(lambda: "Неверный" in pg.inner_text("#auth")); check("old pw rejected", True)
    pg.fill("#pwA","NewPass456!"); pg.click("#authGo"); pg.wait_for_selector("#app:not(.hidden)")
    pg.click("nav >> text=Поиск"); pg.fill("#s_q","Образцова"); pg.wait_for_selector("details.res"); check("data readable after rekey", "Найдено записей: 3" in pg.inner_text("#main"))
    # --- стереть и восстановить
    pg.click("nav >> text=Настройки"); pg.click("text=Стереть все данные"); pg.click(".modal button.danger")
    pg.wait_for_selector("#pwB", timeout=5000); check("wiped -> first run screen", True)
    pg.set_input_files("#restoreIn",bak); pg.wait_for_selector(".modal #rpw"); pg.fill(".modal #rpw",PW); pg.click(".modal button.pri")
    pg.wait_for_selector("text=Данные восстановлены"); pg.click("nav >> text=Отправители")
    check("restore brings senders back", len(pg.query_selector_all("#main table tbody tr"))==5, len(pg.query_selector_all("#main table tbody tr")))
    # --- пустой шаблон и мусор
    pg.click("nav >> text=Файлы"); pg.set_input_files("#fileIn",[D+"bad.xlsx",D+"empty_template.xlsx"]); wait(lambda: len(pg.query_selector_all(".ln"))==2 and "⏳" not in pg.inner_text(".log"))
    L=[x.inner_text() for x in pg.query_selector_all(".ln")]
    check("bad file rejected", L[0].startswith("✖"), L)
    check("empty template accepted with warning", L[1].startswith("✔") and "не заполнен период" in L[1], L)
    # --- сеть
    net=[u for u in reqs if not u.startswith("file:") and not u.startswith("blob:") and not u.startswith("data:")]
    check("no network requests", not net, net)
    r=pg.evaluate("fetch('https://example.com').then(()=>'sent').catch(e=>'blocked')"); check("fetch blocked by CSP", r=="blocked", r)
    check("no JS errors", not errs, errs)
    # --- мобильный вид
    m=br.new_context(viewport={"width":390,"height":780},is_mobile=True,has_touch=True); m.add_init_script(init_script("coord"))
    mp=m.new_page(); mp.goto(URL); mp.screenshot(path=HERE+"/shot_auth_mobile.png")
    pg.click("nav >> text=Сводка"); pg.screenshot(path=HERE+"/shot_home_desktop.png")
    br.close()
print("passed",ok,"failed",len(fails))
sys.exit(1 if fails else 0)
