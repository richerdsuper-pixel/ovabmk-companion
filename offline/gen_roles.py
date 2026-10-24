import openpyxl, glob, json, re, sys, os
SRC = sys.argv[1]
SKIP = {"Yoriqnoma","Boshqaruv","Spravochnik","Hujjatlar","Chop_etish","Eksport"}
RULES = {1:("w",5,12),2:("w",1,12),3:("m",5,17),4:("m",5,17),5:("m",3,12),6:("m",5,17),7:("w",1,12),8:("m",5,12),9:("m",10,17),10:("m",5,12),11:("q",10,17),12:("m",3,12),13:("q",15,17),14:("c",0,0),15:("m",10,17),16:("w",1,12),17:("m",5,17),18:("m",5,17),19:("m",10,17),20:("m",10,17),21:("q",15,17),22:("m",10,17),23:("q",15,17),24:("q",15,17),25:("m",10,17)}
roles=[]
for f in sorted(glob.glob(os.path.join(SRC,"Baza_*.xlsx"))):
    no=int(re.search(r"Baza_(\d\d)_",f).group(1))
    wb=openpyxl.load_workbook(f)
    a1=wb["Yoriqnoma"]["A1"].value
    pos=wb["Boshqaruv"]["C7"].value or ""
    m=re.search(r" / (?=[А-Яа-яЁё])",pos)
    uz,ru=(pos[:m.start()],pos[m.end():]) if m else (pos,pos)
    b=wb["Boshqaruv"]; ind=[]
    for r in range(17,b.max_row+1):
        a=b.cell(r,1).value
        if isinstance(a,(int,float)) and b.cell(r,2).value:
            ind.append({"n":int(a),"row":r,"uz":b.cell(r,2).value,"ru":b.cell(r,3).value,"unit":b.cell(r,4).value,"folder":str(b.cell(r,5).value or "")})
    sheets={}
    for ws in wb:
        if ws.title in SKIP: continue
        hdr=[str(ws.cell(4,c).value or "").split("\n")[-1].strip() for c in range(1,ws.max_column+1)]
        hdr_uz=[str(ws.cell(4,c).value or "").split("\n")[0].strip() for c in range(1,ws.max_column+1)]
        inp=[c-1 for c in range(2,ws.max_column+1) if not str(ws.cell(5,c).value or "").startswith("=")]
        sheets[ws.title]={"title":str(ws["A1"].value or ws.title),"hdr":hdr,"hdrUz":hdr_uz,"inp":inp}
    docs=[]
    if "Hujjatlar" in wb.sheetnames:
        h=wb["Hujjatlar"]
        for r in range(5,h.max_row+1):
            if h.cell(r,3).value: docs.append({"folder":str(h.cell(r,2).value or ""),"uz":h.cell(r,3).value,"ru":h.cell(r,4).value or ""})
    k,d,h=RULES[no]
    roles.append({"no":no,"a1":a1,"uz":uz.strip(),"ru":ru.strip() or uz.strip(),"rule":{"t":k,"d":d,"h":h},"ind":ind,"sheets":sheets,"docs":docs})
json.dump(roles,open("roles.json","w",encoding="utf-8"),ensure_ascii=False,separators=(",",":"))
print(len(roles),os.path.getsize("roles.json"))
