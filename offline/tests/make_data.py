import openpyxl, datetime as dt, sys, os, subprocess, shutil
SRC="/mnt/user-data/uploads/e/OvaBMK/02_Claude_hujjatlari/Excel_dlya_otpravki_25_fayl/"
OUT=os.path.join(os.path.dirname(__file__),"data"); os.makedirs(OUT,exist_ok=True)
tmp=OUT+"/_raw"; os.makedirs(tmp,exist_ok=True)
D=dt.date
NAMES=["Тестова Анна Ивановна","Пробная Мария Петровна","Образцова Зилола Каримовна","Вымышленная Нигора Алиевна","Условная Дилноза Рустамовна","Примерова Гулнора Саидовна","Демонстрационная Лола Акбаровна","Учебная Саида Баховадиновна","Тестова Малика Олимовна","Фиктивная Дильбар Юсуфовна"]
def fill_pregnant(ws,n,base):
    for i in range(n):
        r=5+i
        ws.cell(r,2,NAMES[i%10]); ws.cell(r,3,D(1995+i%8,3+i%6,10+i)); ws.cell(r,4,"3000000000%04d"%i); ws.cell(r,5,"+99890000%04d"%i)
        ws.cell(r,6,"Тестовая махалля"); ws.cell(r,9,base-dt.timedelta(days=40+i*5)); ws.cell(r,10,base-dt.timedelta(days=110+i*5)); ws.cell(r,17,"Ha" if i%3==0 else "Yo'q")
def mk(src,name,inst,filler,phone,ps,pe,fn):
    wb=openpyxl.load_workbook(SRC+src)
    b=wb["Boshqaruv"]; b["C5"]=inst; b["C6"]=filler; b["C8"]=phone; b["C9"]=ps; b["C10"]=pe
    fn(wb,pe)
    p=os.path.join(tmp,name); wb.save(p)
def d1(n):
    def f(wb,pe): fill_pregnant(wb["Homiladorlar"],n,pe)
    return f
def d4(wb,pe):
    ws=wb["Bolalar"]
    for i in range(6):
        r=5+i; ws.cell(r,2,"Детский Тест %d Иванович"%i); ws.cell(r,3,D(2024,1+i,5)); ws.cell(r,4,"3100000000%04d"%i); ws.cell(r,5,"+99891000%04d"%i)
def d2(wb,pe):
    ws=wb["Homiladorlar"]
    for i in range(4):
        r=5+i; ws.cell(r,2,NAMES[i]); ws.cell(r,3,D(1996+i,5,3+i)); ws.cell(r,4,"3200000000%04d"%i); ws.cell(r,5,"+99893000%04d"%i); ws.cell(r,9,pe-dt.timedelta(days=30+i)); ws.cell(r,10,pe-dt.timedelta(days=100+i))
mk("Baza_01_Doya_QVP.xlsx","Baza_01_Doya_QVP_SVP3.xlsx","СВП №3 (тест)","Тестова Дильфуза","+998900000001",D(2026,10,16),D(2026,10,23),d1(10))
mk("Baza_01_Doya_QVP.xlsx","Baza_01_Doya_QVP_SVP5.xlsx","СВП №5 (тест)","Пробная Ойгул","+998900000002",D(2026,10,9),D(2026,10,16),d1(6))
mk("Baza_04_Pediatr.xlsx","Baza_04_Pediatr_SVP3.xlsx","СВП №3 (тест)","Детская Тест Ольга","+998900000003",D(2026,9,1),D(2026,9,30),d4)
mk("Baza_02_Akusher_ginekolog.xlsx","Baza_02_Akusher_RB.xlsx","Район. больница (тест)","Акушерова Тест","+998900000004",D(2026,10,12),D(2026,10,18),d2)
# повторная отправка той же СВП3 за ту же неделю с большим числом записей
mk("Baza_01_Doya_QVP.xlsx","Baza_01_Doya_QVP_SVP3_v2.xlsx","СВП №3 (тест)","Тестова Дильфуза","+998900000001",D(2026,10,16),D(2026,10,23),d1(12))
subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",OUT]+[os.path.join(tmp,f) for f in os.listdir(tmp)],check=True,capture_output=True)
shutil.rmtree(tmp)
print(sorted(os.listdir(OUT)))
